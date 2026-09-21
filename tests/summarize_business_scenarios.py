"""Build the human-readable evidence report from saved scenario results."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import platform
import subprocess
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    out = args.root
    cases = json.loads((out / "final/summary.json").read_text())
    probes = json.loads((out / "probes-verified/boundary.json").read_text())
    baseline = json.loads((out / "baseline/baseline.json").read_text())
    checks = [(s["id"], c) for s in cases + [probes] for c in s["checks"]]
    passed = sum(c["ok"] for _, c in checks)
    extracts = [c for s in cases for c in s["checks"] if c["category"] == "extraction"]
    files = sorted((out / "final").glob("*/inputs/*")) + sorted((out / "probes-verified/90/inputs").glob("*"))
    formats = Counter(p.suffix for p in files)
    # Group repeated observations into distinct issues. A failed check is not a distinct bug.
    issues = [
        ("F01", "高", "识别缺陷", "中文图片金额缩小1000倍", "03",
         "原件1,200.00，首次识别1.200；该样本为中文字体渲染的清晰图片，数字字距较宽。人工修正后完成报表。",
         "先修金额/OCR交叉校验，异常金额不应仅依赖人察觉。", "final/03/result.json"),
        ("F02", "高", "识别及业务缺口", "贷项负号丢失且负数发票无法入账", "90",
         "英文Credit Note总额-120.00识别为+120.00；人工补回负总额、净额和税额后，过账返回HTTP400，提示缺合计金额。",
         "贷项应保留符号，并实现收入/费用与税额同时反向的分录。现有红冲不等同于业务贷项。", "probes-verified/boundary.json"),
        ("F03", "中", "识别缺陷", "购买方名称夹带字段标签", "02/03/05/07/09/12/14",
         "购买方名称被提为“名称”或“名称: 测试企业客户”等，PDF/图片/Word/Excel均有样本。",
         "补全中文标签匹配，避免将“购买方名称”只匹配到“购买方”。", "final/02/result.json"),
        ("F04", "中", "识别缺口", "中文单列金额标签未识别为未税额", "02/03/05/07/09/12/14",
         "只有一个项目的模拟中文单据以“金额”标记未税额，系统留空；并非所有中文税票版式都受影响。",
         "结合版面/项目区区分金额含义，不能无条件把所有“金额”都当小计。", "final/05/result.json"),
        ("F05", "中", "识别缺陷", "USDT币种漏识别", "14",
         "合成发票明确标注Currency: USDT，并带Ethereum地址，首次币种为空；人工补录后估值入账成功。",
         "统一币种代码解析，币种缺失必须继续拦截或提示。", "final/14/result.json"),
        ("F06", "低", "OCR质量", "英文主体名称空格丢失", "08",
         "Cedar Health Services LLC被识别为CedarHealthServicesLLC，影响名称标准化及匹配。",
         "对手方匹配应考虑OCR空格误差，但仍保留原文供审核。", "final/08/result.json"),
        ("F07", "中", "原件下载语义", "中文Word单据下载返回转换PDF", "07",
         "单据original接口返回PDF字节；合集source接口返回真正DOCX且字节完全一致，原件仍保存，没有证据表明原文件丢失。",
         "界面明确区分原始上传件与转换预览件，原件下载应指向source。", "final/07/original-download-diagnosis.json"),
        ("F08", "中", "科目及报表适配边界", "借款暂列其它应付款时现金流勾稽阻断", "12",
         "默认科目表未预设借款科目；用2280其它应付款记1000美元借款并标筹资活动，E4差额1000。系统将该科目余额纳入经营营运资金，阻止导出。",
         "先配置专用借款科目及报表/现金流映射；不能通过把筹资改成经营来掩盖差额。该结果含场景科目选择影响，不是所有融资业务必然失败。", "final/12/reports-after-close.json"),
    ]
    columns = ["编号", "主体（全部虚拟）", "行业", "语言", "币种", "主发票格式", "通过/失败", "人工更正字段数", "报表"]
    table = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    entity_rows = []
    for s in cases:
        row = [s["id"], s["name"], s["sector"], "中文" if s["language"] == "zh" else "英文", s["currency"], s["format"].upper(),
               f'{s["passed"]}/{s["failed"]}', len(s["manual_corrections"]), "已导出" if s["report_exported"] else "E4阻断"]
        entity_rows.append(row)
        linked = row.copy(); linked[1] = f'[{s["name"]}](final/{s["id"]}/result.json)'
        table.append("| " + " | ".join(map(str, linked)) + " |")
    issue_text = []
    for ident, severity, category, title, subjects, detail, next_step, evidence in issues:
        issue_text.append(f"### {ident} · {title}\n\n- 优先级：{severity}；类型：{category}；主体：{subjects}。\n- 实测：{detail}\n- 后续方向：{next_step}\n- [可复核证据]({evidence})。\n")
    metadata = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(), "python": platform.python_version(),
        "production_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "businesses": len(cases), "checks": len(checks), "passed": passed, "failed": len(checks)-passed,
        "files": len(files), "formats": dict(formats),
        "baseline": {k: baseline[k] for k in ("run", "failures", "errors", "skipped")},
        "test_runner_sha256": hashlib.sha256((ROOT / "tests/run_business_scenarios.py").read_bytes()).hexdigest(),
        "inputs": [{"file": str(p.relative_to(out)), "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in files],
    }
    (out / "evidence-manifest.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    report = f"""# 16个中英文主体实测报告

**测试日期：2026-09-18。结论：16个独立虚拟主体均已实际执行测试；15个在必要的人工纠正后完成报表导出，1个融资场景被现金流勾稽拦截。不能把这次结果表述为“16个主体全部自动处理通过”。**

## 实测数量与方法

- 常见行业10个，金融服务3个，加密相关3个；中文7个、英文9个。
- 使用真实FastAPI上传、审核、入账、结算、关账、报表接口，经进程内ASGI TestClient运行；未启动或改动线上服务。
- 共保留{len(files)}份合成源文件（包含两份用于合并PDF的源页），格式分布：{', '.join(f'{k}={v}' for k,v in sorted(formats.items()))}。另有3份内存构造的无效上传内容，不计入源文件数。
- 主体流程检查1067项：1031通过、36未达预期。额外边界检查{len(probes['checks'])}项：{probes['passed']}通过、{probes['failed']}未达预期。合计**{len(checks)}项检查，{passed}通过、{len(checks)-passed}未达预期**。同一缺陷可能影响多个字段/主体，失败项不是独立缺陷数。
- 主体发票首次识别字段核对{len(extracts)}项：{sum(c['ok'] for c in extracts)}准确、{sum(not c['ok'] for c in extracts)}不符。共通过真实审核接口作31处人工字段更正；原始失败保留，没有用更正结果覆盖识别结果。
- 现有回归测试：**{baseline['run']}项，{baseline['failures']}失败、{baseline['errors']}错误、{baseline['skipped']}跳过**。[完整日志](baseline/baseline.log)。新增场景发现了原回归集没有覆盖的边界。
- 每个主体独立SQLite、附件、汇率、导出目录；数据不混账。外部TCP连接在测试进程中被禁止，自动汇率拉取关闭。所有汇率、税额和公司名称均为合成测试值，**不代表实际市场价格、税率或法律适用结论**。
- 原始输入→第一次提取→人工修改→HTTP响应→报表JSON/Excel全程留证。金额与净利润用独立Decimal计算核对，不仅依赖系统自己的借贷平衡。

## 主体明细

{chr(10).join(table)}

每个主体都有销售、采购两份原始单据和一份含4笔交易的流水。中文主体使用中文流水表头。科目是模拟审核人员显式选择的，因此这次**不把人工选科目当作自动分类准确率**。

## 已覆盖的流程

- 文本PDF、图片OCR、Word、Excel、CSV；中英文标签、欧式金额、多币种。
- 期初建账、登记我方主体、上传及原件下载、重复上传、人工修改、版本冲突、审核前禁止入账、重复过账拒绝。
- 应收/应付、两次收款、供应商付款、重复结算拦截；银行流水逐笔日期/金额/余额中，本轮断言具体覆盖金额与余额，日期有保留但未逐字段评分。
- 零售/制造/贸易：库存及部分成本结转；制造/医疗：设备购置和折旧；教育：薪酬；物业：预收款；资管/经纪：代收本金挂负债再转出，不混入收入。
- 融资服务：借款本金、利息支出与筹资现金流分类，发现默认科目适配限制。
- 16个主体的销售折算金额与净利润均与独立预期一致（发生在人工更正后）；15个主体关账前后五项勾稽通过并导出7张工作表，发票号可追溯。
- 额外边界：负数贷项、总额与税额不一致、缺币种、错误钱包格式、PDF多发票拆分、不同汇率的EUR结算、ETH原币高精度结算、自动对账→人工确认→据匹配结算、空文件/不支持类型/损坏PDF、隔离账套备份与一致性核验。

## 发现的问题

{chr(10).join(issue_text)}

## 加密与金融业务的实测边界

- USDT：首次识别漏币种；人工补录后使用0.997的合成估值，没有默认按1美元处理。无汇率入账被拒绝。
- ETH：`1.23456789 × 2400 = 2962.96 USD`（本位币舍入至分）核对通过，原币结算接口也核对通过。
- BTC：`0.12345678 × 60000 = 7407.41 USD`核对通过；钱包字符串仅作格式/风险测试，没有任何链上支付或余额查询。
- **当前HTTP结算将ETH收款记入1002银行存款。** 这验证了原币数值换算，不证明系统具备钱包余额、币种数量台账、交易哈希核验、Gas归集、质押奖励自动入账或代币估值损益模型。直接用于真实加密资产核算仍需单独设计。
- 资管/经纪测试覆盖服务费及代收待付款，不覆盖基金NAV、持仓估值、客户资产隔离账簿、证券交割、衍生品保证金、利息摊销或金融工具减值。这些不能由普通服务发票测试推导为已支持。
- 税额均为测试输入，只检查提取、算术和分录；没有验证各国税务申报或准则适用。

## 证据与复跑

- [测试明细Excel](测试明细.xlsx)：主体、逐项检查、问题汇总、首次提取、合成汇率五张表。
- [主体结果JSON](final/summary.json)、[边界结果JSON](probes-verified/boundary.json)、[版本和输入哈希清单](evidence-manifest.json)。
- 每个`final/编号/`下有`inputs/`源文件、`result.json`请求响应与修改记录、关账前后报表JSON；成功者另有`financial-statements.xlsx`。测试数据库保存在各自`runtime/`中。
- [合成样例与测试结果压缩包](合成样例与测试结果.zip)：仅包含本轮合成输入、结果、报告和复跑脚本，排除runtime目录及早期调试轮次。

在项目根目录、已有依赖的环境中执行（请使用新的输出目录）：

```bash
.venv/bin/python tests/run_business_scenarios.py --output data/scenario-retest/businesses
.venv/bin/python tests/run_business_scenarios.py --output data/scenario-retest/boundaries --probes
.venv/bin/python tests/run_business_scenarios.py --output data/scenario-retest/baseline --baseline
```

脚本会执行完全部主体并保存结果；存在未达预期项时退出码为1，基线成功为0。选测可加`--ids 03,14`。不覆盖已有主体目录。

本轮为**接口和业务流程实测**，页面只检查可访问，未进行浏览器点击/布局自动化；不包含多人同步、权限改造、压力测试或生产部署。生产业务代码保持原样，本轮新增测试和证据记录。
"""
    (out / "测试报告.md").write_text(report, encoding="utf-8")
    w = Workbook(); w.remove(w.active)
    sheets = {
        "主体总览": [columns] + entity_rows,
        "逐项检查": [["主体编号", "分类", "检查", "结果", "实际", "期望"]] +
                      [[ident, c["category"], c["name"], "通过" if c["ok"] else "未达预期",
                        json.dumps(c.get("actual"), ensure_ascii=False), json.dumps(c.get("expected"), ensure_ascii=False)] for ident, c in checks],
        "问题汇总": [["编号", "优先级", "类型", "问题", "主体", "实测", "后续方向", "证据"]] + [list(x) for x in issues],
        "首次提取": [["主体", "文件", "字段", "期望", "首次提取", "识别方式"]] +
                    [[s["id"], d["file"], k, str(v), str(d["extracted"].get(k)), d.get("parse_method", "")]
                     for s in cases for d in s["documents"] for k,v in d["truth"].items()],
        "合成汇率": [["币种", "测试USD估值（非市场价）"]] + list(json.loads((out / "final/manifest.json").read_text())["rates"].items()),
    }
    for name, rows in sheets.items():
        sh = w.create_sheet(name)
        for row in rows:
            sh.append(row)
        sh.freeze_panes = "A2"; sh.auto_filter.ref = sh.dimensions
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17365D")
        for col in sh.columns:
            letter = col[0].column_letter
            sh.column_dimensions[letter].width = min(64, max(16, max(len(str(c.value or "")) for c in col[:80]) + 2))
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if name == "逐项检查":
            for row in sh.iter_rows(min_row=2):
                if row[3].value == "未达预期":
                    row[3].fill = PatternFill("solid", fgColor="FCE4D6")
    w.save(out / "测试明细.xlsx"); w.close()
    (out / "issues.json").write_text(json.dumps(issues, ensure_ascii=False, indent=2), encoding="utf-8")
    zip_path = out / "合成样例与测试结果.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in ("测试报告.md", "测试明细.xlsx", "issues.json", "evidence-manifest.json", "baseline/baseline.json", "baseline/baseline.log"):
            z.write(out / name, name)
        for dirname in ("final", "probes-verified"):
            for path in sorted((out / dirname).rglob("*")):
                rel = path.relative_to(out)
                if path.is_file() and "runtime" not in rel.parts:
                    z.write(path, str(rel))
        for name in ("run_business_scenarios.py", "summarize_business_scenarios.py"):
            z.write(ROOT / "tests" / name, f"复跑脚本/{name}")
    print(json.dumps({"report": str(out / "测试报告.md"), "checks": len(checks), "passed": passed,
                      "failed": len(checks)-passed, "zip_bytes": zip_path.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
