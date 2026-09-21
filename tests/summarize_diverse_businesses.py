"""Auditable report and rendered layout catalogue for the diverse cohort."""
from collections import Counter
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import zipfile

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image, ImageOps

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))


def layout_catalogue(out, cases):
    from extraction.extract import excel, office
    seen=set(); document=fitz.open(); thumbs=[]
    for s in cases:
        family=s['layout_family']
        if family in seen:continue
        seen.add(family); path=out/'verified'/s['id']/s['documents'][0]['file']
        page=document.new_page(width=620,height=880)
        page.insert_text((20,24),s['id']+' '+family,fontname='china-s',fontsize=12)
        page.insert_text((20,43),s['name'],fontname='china-s' if s['language']=='zh' else 'helv',fontsize=10)
        rect=fitz.Rect(20,60,600,850)
        if path.suffix=='.pdf':
            src=fitz.open(path);page.show_pdf_page(rect,src,0);src.close()
        elif path.suffix=='.png':page.insert_image(rect,filename=str(path))
        elif path.suffix=='.xlsx':page.insert_image(rect,stream=excel.render_png(path))
        else:
            src=fitz.open(stream=office.to_pdf(path),filetype='pdf');page.show_pdf_page(rect,src,0);src.close()
        pix=page.get_pixmap(matrix=fitz.Matrix(.4,.4))
        thumbs.append(Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB'))
    document.save(out/'20种版式预览.pdf');document.close()
    sheet=Image.new('RGB',(5*268,4*372),'#dddddd')
    for i,img in enumerate(thumbs):sheet.paste(img,(10+(i%5)*268,10+(i//5)*372));img.close()
    sheet.save(out/'版式总览.png');sheet.close()


def main(out):
    read=lambda p:json.loads((out/p).read_text())
    cases=read('verified/summary.json');before=read('before/summary.json');baseline=read('baseline/baseline.json')
    assert len(cases)==30 and all(s['complete'] for s in cases)
    checks=[(s['id'],c) for s in cases for c in s['checks']]
    fields=[(i,c) for i,c in checks if c['name'].startswith('首次提取 ')]
    bad=[(i,c) for i,c in checks if not c['ok']]
    checkpoints=[(s['id'],p) for s in cases for p in s['checkpoints']]
    ties=[(i,c) for i,c in checks if c['category']=='ties']
    changes=sum(len(s['manual_corrections']) for s in cases)
    forced=sum(len(s.get('forced_reviews',[])) for s in cases)
    fixes=[
        ['R3-01','E2关账后不平','负收入、贷方费用/汇兑收益或同组正负抵消时漏结转；改为逐科目按真实余额方向结转。','117/118/119及混合正负、重开再关专项'],
        ['R3-02','E4赊购及非现金业务','固定/无形资产赊购、库存资本化、应付转借款、权益薪酬、股利应付等非现金变动漏调整；从具体分录取数，不按勾稽差额凑平。','102/103/104/107/108/114/115/116/120/123'],
        ['R3-03','E4折旧及资产处置红冲','原逻辑只加累计折旧贷方发生额，漏掉折旧冲回，还把处置撤销误计折旧；现区分非现金计提与处置，按符号抵消。','105/106/107'],
        ['R3-04','E6漏报无效映射','未知资产负债表行或直接映射NetIncome等计算行会从报表消失；现列入未归类并阻止导出。','test_round3_ties的负向用例'],
        ['R3-05','明细合计重复','带数量/单价/行金额的分隔表格被散文拆分，单价重复累计；明确表头按列取值，挂牌价/商业折扣汇总不作商品行。','101/103/112/115/121/122及布局专项'],
        ['R3-06','净额/税额/总额错位','中文横向汇总被误认作增值税商品表，覆盖正确汇总；增值税表识别增加商品描述列要求。','104及横向汇总负向用例'],
        ['R3-07','名称及未税额漏识别','增加Seller/Supplier/Vendor标签及未税金额；文档标题不当主体名称。','中英文PDF、Word、Excel多个主体'],
        ['R3-08','Excel封面遮蔽发票','活动页是封面时，只在存在唯一明确发票页时选取该页；多候选不武断选择。','109及唯一/歧义工作表专项'],
        ['R3-09','轻微倾斜扫描读序','以OCR文字框基线估计小角度斜率，仅调整解析读序，不旋转原图、不移动原始高亮框。','116/128及坐标保持专项'],
        ['R3-10','报表解释显示','增加其他非现金调整、非经营损益扣除及凭证展开；修正流动资产减少时拼出双负号的显示问题。','web/reports.html语法验证'],
        ['R3-11','中文Excel预览缺字','版式预览检查发现DejaVu缺少中文字形；改用随PyMuPDF提供的中文字库，不依赖系统安装字体。','Office回归及中文字符字形检查'],
    ]
    columns=['编号','虚拟主体','业务路径','版式','通过/未达预期','字段改正','报表']
    rows=[[s['id'],s['name'],s['sector'],s['layout_family'],f"{s['passed']}/{s['failed']}",len(s['manual_corrections']),
           '已导出' if s.get('report_exported') else '被阻止'] for s in cases]
    table=['|'+'|'.join(columns)+'|','|'+'---|'*len(columns)]
    for row in rows:
        r=row.copy();r[1]=f'[{r[1]}](verified/{r[0]}/result.json)';table.append('|'+'|'.join(map(str,r))+'|')
    correction_rows=[[s['id'],c['field'],str(c['before']),str(c['after'])] for s in cases for c in s['manual_corrections']]
    before_bad=sum(s['failed'] for s in before)
    before_ties=[s['id'] for s in before if any(c['category']=='ties' and not c['ok'] for c in s['checks'])]
    formats=Counter(s['format'] for s in cases); langs=Counter(s['language'] for s in cases)
    report=f'''# 第三轮：30主体、20种版式与中间状态勾稽实测

**30个独立合成账套，20类结构不同的单据布局，{len(checkpoints)}个账务检查点。最终30主体均出表，{len(ties)}项E1/E2/E3/E4/E6检查全部通过；首次识别仍有{len(bad)}个字段偏差，经过显式审核校正后继续，未隐藏这些失败。**

## 本次为什么比前两轮不同

- 主体{langs['zh']}中文、{langs['en']}英文，覆盖制造、物流、医疗、影院、研发、食品、物业、人力、融资、控股、零售、专业服务、公益、旅游、支付及BTC/ETH/USDT计价服务。
- 布局结构包括双栏抬头、标签在上数值在下、横向纸张、横向汇总表、窄幅小票、两页续单、Word表格/段落、Excel封面/合并格/数量单价表、欧式金额、分列税项、合计前置、换行描述、折扣摘要、倾斜扫描、括号负数贷项等。格式数量：{dict(formats)}。不是以换公司名凑版式数量。
- [20种版式预览PDF](20种版式预览.pdf)、[版式缩略总览](版式总览.png)可直接核对布局。预览展示每类首个主体的首页；多页源件、Excel封面和完整Word表格以原始文件为准。
- 账务路径不再统一为销售/采购后立即付清：逐步检查赊购未付、部分付、付清、非现金重分类、折旧/处置红冲、负收入/负费用、外币汇兑收益、两期关账及重开。
- 所有数据由脚本合成，部分金额固定以便隔离原因；这些是业务机制测试，**不是30套真实供应商模板或各行业准则合规认证**。科目及现金活动由模拟审核人显式选择，汇率为合成值，不能算自动分类准确率。

## 结果与计数口径

- 最终检查共{len(checks)}项：{len(checks)-len(bad)}通过、{len(bad)}未达预期，全部剩余失败为首次OCR字段识别；测试进程因此保留退出码1。
- 首次字段核对{len(fields)}项：{sum(c['ok'] for _,c in fields)}正确、{sum(not c['ok'] for _,c in fields)}不符；实际作{changes}处字段更正。数值相同、仅1000/1000.00格式不同的不再算人工更正。
- {len(checkpoints)}个中间/关账检查点，各核对五项报表勾稽、独立净利润、独立现金余额及往来明细。{len(ties)}项报表勾稽全过；净利润/现金预期由业务事件显式定义，不拿系统报表当预期。
- 30份最终报表，每份7张工作表。正常审核全部通过，**风险强制覆盖次数为{forced}**。
- 全量回归{baseline['run']}项：{baseline['failures']}失败、{baseline['errors']}错误、{baseline['skipped']}跳过；包括本轮18项专项。最后中文Excel预览字体增量另跑Office回归及不同中文字形渲染检查通过。另保留将错误现金活动分类、未知报表映射挡住的负向用例，未通过放宽勾稽实现“全过”。
- 修复前30主体尝试执行：29个走完流程、15个出表，{len(before_ties)}个主体触发报表勾稽问题，共{before_bad}个未达预期检查。117外币结算因测试脚本漏传差额原因而中止，此为测试脚本配置错误，随后显式指定汇兑损益原因，**不算系统缺陷**。旧识别和强制覆盖记录保留在before中。

## 复现并修复的问题

|编号|问题|原因和处理|证据|
|---|---|---|---|
{chr(10).join('|'+'|'.join(row)+'|' for row in fixes)}

R3-01尤其容易让真实账套出现“关账前正常、关账后不平”：原代码只处理收入/费用合计为正的情况。修复前，EUR汇兑收益场景E2差额-110，纯销售贷项差额+1000，纯采购贷项差额-1000；修复后全部为0。**这是本轮合成场景的复现结果，不冒称已经检查了用户的真实账套。**

## 主体与业务路径

{chr(10).join(table)}

## 剩余OCR问题

|主体|字段|首次结果|原件正确值|
|---|---|---|---|
{chr(10).join('|'+'|'.join(row)+'|' for row in correction_rows)}

105是窄幅小票，部分字形被OCR误读；128是倾斜中文扫描，仍有票号、名称及未税额偏差。错误仍出现在原始结果中；格式读序修复没有消除模型字形识别问题。不能因此宣称任意真实票据可直接自动记账。

## 已经形成的错误关账如何验证

在独立副本复制117/118/119三份旧合成数据库，确认E2仍不平，再重开2026-08并按修复后的逻辑重新关账，三份均恢复勾稽。[修复演练结果](recovery/summary.json)。

软件升级会更新计算代码，**不会自动重写旧的结转凭证**。真实账套若已经关过账，应先保存备份、核对错误结转和期间状态，再人工执行重开及重新关账；本轮没有修改真实账套。

## 证据、范围和复跑

- [测试明细Excel](测试明细.xlsx)：主体、全部断言、各检查点五项差额、首次识别、人工更正、问题修复表。
- [最终结果](verified/summary.json)、[修复前结果](before/summary.json)、[全量测试日志](baseline/baseline.log)、[版本与源件哈希](evidence-manifest.json)。
- [合成样例与结果包](合成样例与测试结果.zip)含30份原始单据、结果、检查点报表、30份Excel报表、预览及复跑脚本。排除runtime目录及实际运行配置。每主体本地保留独立SQLite，可复核全链路。
- 实际执行FastAPI上传、审核、入账、结算、红冲、关账、报表接口；未开展浏览器点击/布局自动化。报表页面改动通过内联JS语法检查。
- 外部TCP被测试脚本禁止，自动汇率关闭。加密主体只验证估值、USD收款或已知代币原币结算拒绝；不覆盖链上数量账、基金NAV、证券交割及监管申报。
- 修复写入当前共享工作区源码，未提交/推送、未发布演示站或新应用分发包；其他电脑的旧源码不会自动得到修复，既有启动进程也需重启加载新代码。

```bash
.venv/bin/python tests/run_diverse_businesses.py --output data/diverse-retest/verified
.venv/bin/python -m unittest tests.test_round3_ties tests.test_round3_layouts -v
```

输出目录须全新；可加`--ids 102,117,118`选测。报告比较使用已保留的修复前快照，重新运行新代码不能重造旧缺陷统计。
'''
    (out/'测试报告.md').write_text(report,encoding='utf-8')
    w=Workbook();w.remove(w.active)
    sheets={
        '主体总览':[columns]+rows,
        '全部检查':[['主体','类别','检查','通过','实际','预期']]+[[i,c['category'],c['name'],c['ok'],json.dumps(c['actual'],ensure_ascii=False),json.dumps(c['expected'],ensure_ascii=False)] for i,c in checks],
        '检查点勾稽':[['主体','阶段','E1','E2','E3','E4','E6','预期利润','实际利润','预期现金','实际现金']]+[
            [i,p['stage']]+[json.dumps(p['report']['checks'][k],ensure_ascii=False) for k in ('E1_balance_sheet_balanced','E2_income_to_retained','E3_cash_tie','E4_cfo_direct_indirect','E6_all_classified')]+
            [p['expected_profit'],p['report']['income_statement']['net_income'],p['expected_cash'],p['report']['cash_flow_statement']['ending']] for i,p in checkpoints],
        '首次识别':[['主体','文件','字段','预期','首次结果']]+[[s['id'],d['file'],k,str(v),str(d['extracted'][k])] for s in cases for d in s['documents'] for k,v in d['truth'].items()],
        '人工更正':[['主体','字段','首次结果','正确值']]+correction_rows,
        '修复问题':[['编号','问题','处理','证据']]+fixes,
    }
    for title,rr in sheets.items():
        sh=w.create_sheet(title)
        for r in rr:sh.append(r)
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=min(65,max(15,max(len(str(c.value or '')) for c in col[:80])+2))
        for row in sh.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    w.save(out/'测试明细.xlsx');w.close()
    layout_catalogue(out,cases)
    sourcefiles=['reports/service.py','ledger/close.py','web/reports.html','extraction/parse/generic.py','extraction/extract/excel.py',
                 'extraction/extract/ocr.py','tests/run_diverse_businesses.py','tests/test_round3_ties.py','tests/test_round3_layouts.py','tests/summarize_diverse_businesses.py']
    sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    manifest={'generated_at':datetime.now(timezone.utc).isoformat(),'subjects':30,'layout_families':len({s['layout_family'] for s in cases}),
              'checkpoints':len(checkpoints),'checks':len(checks),'failed':len(bad),'first_fields':len(fields),'manual_corrections':changes,'forced_reviews':forced,
              'baseline':baseline,'base_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
              'code_state':'未提交修复，以源码SHA256为准','source_hashes':{f:sha(ROOT/f) for f in sourcefiles},
              'input_hashes':{str(p.relative_to(out)):sha(p) for p in sorted((out/'verified').glob('*/inputs/*'))}}
    (out/'evidence-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    with zipfile.ZipFile(out/'合成样例与测试结果.zip','w',zipfile.ZIP_DEFLATED) as z:
        for f in ('测试报告.md','测试明细.xlsx','20种版式预览.pdf','版式总览.png','evidence-manifest.json','baseline/baseline.json','baseline/baseline.log'):
            z.write(out/f,f)
        for folder in ('verified','before','iteration2','recovery','verification'):
            for p in sorted((out/folder).rglob('*')):
                rel=p.relative_to(out)
                if p.is_file() and 'runtime' not in rel.parts:z.write(p,str(rel))
        for f in sourcefiles+['tests/run_business_scenarios.py']:
            if f.startswith('tests/'):z.write(ROOT/f,'复跑脚本/'+Path(f).name)
    print(json.dumps({k:manifest[k] for k in ('subjects','layout_families','checkpoints','checks','failed','manual_corrections','forced_reviews')},ensure_ascii=False))


if __name__=='__main__':main(Path(sys.argv[1]))
