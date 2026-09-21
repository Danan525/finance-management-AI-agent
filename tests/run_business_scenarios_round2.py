"""Second independent cohort; reuse the first-round HTTP harness, never its data."""
from __future__ import annotations

import argparse
import io
import json
import logging
from pathlib import Path
import traceback
from decimal import Decimal as D

import run_business_scenarios as base
from core import db
from ledger import store
from openpyxl import load_workbook

BUSINESSES = [
    ("21", "Maple Grid Software LLC", "en", "软件订阅", "USD", "pdf", "subscription"),
    ("22", "苏州禾澄商贸有限公司", "zh", "零售电商", "CNY", "xlsx", "inventory"),
    ("23", "成都竹里餐饮有限公司", "zh", "餐饮扫描票据", "CNY", "png", "restaurant"),
    ("24", "Westfield Precision GmbH", "en", "制造与欧式金额", "EUR", "pdf", "manufacturing"),
    ("25", "武汉江桥供应链有限公司", "zh", "物流运输", "CNY", "pdf", "logistics"),
    ("26", "Riverbend Creative Ltd", "en", "专业设计", "GBP", "docx", "design"),
    ("27", "厦门知远培训有限公司", "zh", "教育与Word单据", "CNY", "docx", "education"),
    ("28", "Willow Health Services LLC", "en", "医疗与英文OCR", "USD", "png", "health"),
    ("29", "成都栖山物业服务有限公司", "zh", "物业与预收款", "CNY", "xlsx", "property"),
    ("30", "Crestline Pacific Trading Ltd", "en", "跨境贸易", "HKD", "pdf", "trading"),
    ("31", "Juniper Fund Services Ltd", "en", "金融/资管服务", "USD", "pdf", "fund"),
    ("32", "杭州澄明融资服务有限公司", "zh", "金融/借款与利息", "CNY", "xlsx", "lending"),
    ("33", "Granite Brokerage Services Ltd", "en", "金融/经纪服务", "USD", "xlsx", "broker"),
    ("34", "苏州星海区块链技术有限公司", "zh", "加密/USDT计价服务", "USDT", "pdf", "stablecoin"),
    ("35", "Lumen Validator Services Ltd", "en", "加密/ETH计价服务", "ETH", "pdf", "validator"),
    ("36", "Harbor Digital Custody Ltd", "en", "加密/BTC计价服务", "BTC", "xlsx", "treasury"),
    ("37", "Bluewater Retail Returns Ltd", "en", "零售及销售贷项退款", "USD", "pdf", "inventory"),
    ("38", "宁波锦云设计有限公司", "zh", "设计及采购贷项退款", "CNY", "pdf", "design"),
    ("39", "Marsh Telecom Services Ltd", "en", "通信订阅", "EUR", "xlsx", "subscription"),
    ("40", "成都竹林酒店服务有限公司", "zh", "酒店服务及预收款", "CNY", "pdf", "property"),
    ("41", "Canopy Renewable Engineering LLC", "en", "新能源工程与设备", "USD", "docx", "manufacturing"),
    ("42", "苏州远星文化传媒有限公司", "zh", "文化传媒图片票据", "CNY", "png", "design"),
    ("43", "Meadow Learning Foundation", "en", "非营利培训服务", "USD", "xlsx", "education"),
    ("44", "Atlas Payment Services Ltd", "en", "金融/支付服务", "USD", "png", "broker"),
]
REPEATS = {
    "22": ["F03", "F04"], "23": ["F01", "F03", "F04"], "25": ["F03", "F04"],
    "27": ["F03", "F04", "F07"], "28": ["F06"], "32": ["F08"],
    "34": ["F05"], "37": ["F02"], "38": ["F02", "F03", "F04"], "42": ["F01", "F03", "F04"],
}
original_rows = base.invoice_rows


def cohort_rows(b, kind="sale", **kw):
    rows, truth = original_rows(b, kind, **kw)
    # Vary amounts independently of the implementation, not just business names.
    if kind in ("sale", "purchase") and truth["currency_settlement"] not in ("BTC", "ETH", "USDT"):
        gross = (D("1200") + D(b[0])*D("17.35")) if kind == "sale" else (D("300") + D(b[0])*D("3.20"))
        tax = (gross * D("0.10")).quantize(D("0.01"))
        net = gross - tax
        changes = {"Total Due": gross, "价税合计": gross, "Tax": tax, "税额": tax, "Subtotal": net, "金额": net}
        rows = [[r[0], base.money(changes[r[0]], euro=(b[4] == "EUR"))] if r[0] in changes else r for r in rows]
        truth.update(total_due=str(gross), sales_tax=str(tax), subtotal=str(net))
    if b[0] == "38" and kind == "credit":
        issuer = "测试办公供应商"
        old_name, old_customer = truth["issuer_name"], truth["customer_name"]
        rows = [[issuer] if len(r) == 1 and r[0] == old_name else
                [r[0], issuer] if len(r) == 2 and r[0] == "销售方名称" else
                [r[0], b[1]] if len(r) == 2 and r[0] == "购买方名称" else r for r in rows]
        truth.update(issuer_name=issuer, customer_name=b[1])
    return rows, truth


base.invoice_rows = cohort_rows


class CohortScenario(base.Scenario):
    def __init__(self, b, root):
        super().__init__(b, root)
        self.result["previous_problem_targets"] = REPEATS.get(b[0], [])

    def extract(self, *args, **kwargs):
        h, truth, path = super().extract(*args, **kwargs)
        inv = db.get_invoice(h)
        self.result["documents"][-1]["field_provenance"] = {k: {"raw": inv.f(k).raw, "source": inv.f(k).source,
            "note": inv.f(k).note} for k in truth}
        return h, truth, path

    def final_report(self, expected_profit):
        _, report = self.call("GET", "/api/reports")
        self.check("扩展业务后净利润独立核对", D(report["income_statement"]["net_income"]) == expected_profit,
                   report["income_statement"]["net_income"], str(expected_profit), "accounting")
        self.post("/api/ledger/close", {"period": "2026-08"}, "扩展业务后重新关账")
        _, report = self.call("GET", "/api/reports")
        self.check("扩展业务后报表五项勾稽", report["can_issue"], report["checks"], category="reports")
        exp = self.post("/api/reports/export", {}, "导出最终报表")
        r = self.client.get(exp["download"])
        w = load_workbook(io.BytesIO(r.content), read_only=True)
        self.check("最终Excel七张表", len(w.sheetnames) == 7, w.sheetnames, 7, "reports"); w.close()
        (self.root / "financial-statements.xlsx").write_bytes(r.content)
        base.dump(self.root / "reports-after-close.json", report)
        self.result["expected_profit_usd"] = str(expected_profit)
        self.result["income_statement"] = report["income_statement"]

    def run(self):
        super().run()
        self.result["complete"] = False
        no = self.b[0]
        expected = D(self.result["expected_profit_usd"])
        if no in ("37", "38"):
            self.post("/api/ledger/reopen", {"period": "2026-08"}, "重开测试期间追加退货贷项")
            h, truth, _ = self.extract("credit", fmt="pdf")
            self.correct(h, truth)
            direction = "AR" if no == "37" else "AP"
            self.approve(h, "4000" if direction == "AR" else "6900")
            entry = self.post("/api/ledger/post", {"file_hash": h, "direction": direction, "as_of": "2026-08-05"}, "贷项过账")
            je = store.get_entry(entry["entry_no"])
            tax_code = "2210" if direction == "AR" else "1180"
            tax_line = next(l for l in je.lines if l.account.startswith(tax_code))
            rate = D(base.RATES[truth["currency_settlement"]])
            gross = (abs(D(truth["total_due"])) * rate).quantize(D("0.01"))
            net = (abs(D(truth["subtotal"])) * rate).quantize(D("0.01"))
            self.check("贷项税额反向且金额正确", (tax_line.debit if direction == "AR" else tax_line.credit) == gross-net,
                       vars(tax_line), str(gross-net), "accounting")
            first = (gross / 3).quantize(D("0.01"))
            refund = self.post("/api/ledger/settle", {"file_hash": h, "cash_amount": str(first), "settle_amount": str(first),
                                "date": "2026-08-21"}, "贷项第一笔退款")
            delta = store.get_entry(refund["entry_no"]).cash_delta()
            self.check("退款现金方向正确", delta == (-first if direction == "AR" else first), str(delta), category="accounting")
            self.post("/api/ledger/settle", {"file_hash": h, "cash_amount": str(gross-first), "date": "2026-08-22"}, "贷项退款结清")
            _, view = self.call("GET", "/api/ledger/open")
            self.check("贷项结清后往来控制与明细一致", all(v["ok"] for v in view["control"].values()), view["control"], category="accounting")
            expected += -net if direction == "AR" else net
            self.final_report(expected)
        if no in ("34", "35", "36"):
            self.post("/api/ledger/reopen", {"period": "2026-08"}, "重开测试期间验证原币结算边界")
            h, truth, _ = self.extract("digital-asset", fmt="pdf")
            self.correct(h, truth); self.approve(h, "4000")
            self.post("/api/ledger/post", {"file_hash": h, "direction": "AR", "as_of": "2026-08-05"}, "第二份代币计价服务应收入账")
            before = self.call("GET", "/api/ledger/entries")[1]
            valuation = (D(truth["total_due"]) * D(base.RATES[truth["currency_settlement"]])).quantize(D("0.01"))
            status, data = self.call("POST", "/api/ledger/settle", {"file_hash": h, "cash_amount": truth["total_due"],
                "cash_currency": truth["currency_settlement"], "settle_amount": str(valuation), "date": "2026-08-25"})
            self.check("代币原币收付不误记银行", status == 400 and "数字资产" in str(data), data, category="crypto-safeguard")
            self.check("拒绝代币结算未新增凭证", self.call("GET", "/api/ledger/entries")[1] == before, category="crypto-safeguard")
            expected += valuation
            self.result["crypto_scope"] = "首份合同模拟实际USD收款；第二份保持未结。原币结算被拦截，不代表链上资产子账已实现。"
            self.final_report(expected)
        self.result["complete"] = True


def main():
    p = argparse.ArgumentParser(); p.add_argument("--output", type=Path, required=True)
    p.add_argument("--ids", default=""); p.add_argument("--baseline", action="store_true")
    args = p.parse_args(); args.output.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.ERROR)
    if args.baseline:
        raise SystemExit(0 if base.baseline(args.output) else 1)
    base.dump(args.output / "manifest.json", {"businesses": BUSINESSES, "previous_problem_targets": REPEATS,
               "rates": base.RATES, "synthetic_only": True, "round": 2})
    results = []
    for b in BUSINESSES:
        if args.ids and b[0] not in args.ids.split(","):
            continue
        print(f"START {b[0]} {b[1]}", flush=True); s = None
        try:
            s = CohortScenario(b, args.output); s.run()
        except Exception:
            detail = traceback.format_exc(); print(detail, flush=True)
            if s:
                s.result["fatal"] = detail; s.result["complete"] = False
            else:
                results.append({"id": b[0], "name": b[1], "complete": False, "failed": 1, "fatal": detail})
        finally:
            if s:
                results.append(s.finish())
                print(f"END {b[0]} passed={s.result['passed']} failed={s.result['failed']} complete={s.result['complete']}", flush=True)
        base.dump(args.output / "summary.json", results)
    raise SystemExit(1 if any(r["failed"] or not r["complete"] for r in results) else 0)


if __name__ == "__main__":
    main()
