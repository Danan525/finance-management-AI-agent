"""Generate synthetic businesses and exercise the real HTTP application in isolated ledgers.

Run: .venv/bin/python tests/run_business_scenarios.py --output <new-directory>
All names, documents, amounts and FX rates are fictional test fixtures. No market
prices, tax advice, blockchain transactions or production data are used.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import socket
import sys
import traceback
from decimal import Decimal as D

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("OCR_ENGINE", "rapid")
os.environ.setdefault("OMP_NUM_THREADS", "2")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "2")

# TestClient uses in-process ASGI transport; block all real TCP connections.
def offline_connect(*args, **kwargs):
    raise RuntimeError("Scenario tests prohibit external network connections")

socket.socket.connect = offline_connect
socket.create_connection = offline_connect

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient
from core import config, db

BUSINESSES = [
    ("01", "Northstar Cloud Labs LLC", "en", "软件订阅/SaaS", "USD", "pdf", "subscription"),
    ("02", "广州晴岚零售有限公司", "zh", "零售电商", "CNY", "xlsx", "inventory"),
    ("03", "杭州青禾餐饮有限公司", "zh", "餐饮", "CNY", "png", "restaurant"),
    ("04", "Meridian Components GmbH", "en", "制造业/跨境采购", "EUR", "pdf", "manufacturing"),
    ("05", "深圳远帆供应链有限公司", "zh", "物流运输", "CNY", "pdf", "logistics"),
    ("06", "Harborview Design Studio Ltd", "en", "专业设计服务", "GBP", "docx", "design"),
    ("07", "南京知行培训有限公司", "zh", "教育培训", "CNY", "docx", "education"),
    ("08", "Cedar Health Services LLC", "en", "医疗健康服务", "USD", "png", "health"),
    ("09", "上海栖居物业服务有限公司", "zh", "物业租赁服务", "CNY", "xlsx", "property"),
    ("10", "Pacific Market Trading Ltd", "en", "进出口贸易", "HKD", "pdf", "trading"),
    ("11", "Blue Lantern Asset Management Ltd", "en", "金融/资产管理", "USD", "pdf", "fund"),
    ("12", "上海澄川融资咨询有限公司", "zh", "金融/融资服务", "CNY", "xlsx", "lending"),
    ("13", "Summit Brokerage Services Ltd", "en", "金融/经纪服务", "USD", "xlsx", "broker"),
    ("14", "星桥链上技术有限公司", "zh", "加密/USDT 服务结算", "USDT", "pdf", "stablecoin"),
    ("15", "Aurora Validator Services Ltd", "en", "加密/ETH 验证服务", "ETH", "pdf", "validator"),
    ("16", "Seabrook Digital Treasury Ltd", "en", "加密/BTC 数字资产服务", "BTC", "xlsx", "treasury"),
]
RATES = {"USD": "1", "CNY": "0.14", "EUR": "1.10", "GBP": "1.25",
         "HKD": "0.128", "USDT": "0.997", "ETH": "2400", "BTC": "60000"}
DESCRIPTIONS = {
    "subscription": "Annual cloud software subscription",
    "inventory": "日用品销售及仓储服务", "restaurant": "团体餐饮服务",
    "manufacturing": "Precision component manufacturing", "logistics": "跨境货运服务",
    "design": "Brand design consulting", "education": "企业技能培训服务",
    "health": "Health screening service", "property": "物业管理服务",
    "trading": "Export handling service", "fund": "Quarterly management fee",
    "lending": "融资咨询服务费（不含贷款本金）", "broker": "Brokerage commission, excluding client principal",
    "stablecoin": "区块链技术服务（USDT结算）", "validator": "Validator infrastructure service fee",
    "treasury": "Digital asset custody administration fee",
}


def dump(path, obj):
    Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def money(v, euro=False):
    text = f"{D(v):,.2f}" if D(v).as_tuple().exponent >= -2 else str(v)
    return text.replace(",", "@").replace(".", ",").replace("@", ".") if euro else text


def invoice_rows(b, kind="sale", *, bad_total=False, bad_wallet=False):
    no, name, lang, sector, ccy, fmt, variant = b
    sale = kind != "purchase"
    total = ("0.12345678" if ccy == "BTC" else "1.23456789" if ccy == "ETH" else "1200.00") if sale else "300.00"
    # Purchase is a USD operating expense for crypto businesses.
    actual_ccy = ccy if sale else ("USD" if ccy in ("USDT", "ETH", "BTC") else ccy)
    tax = "0.00" if actual_ccy in ("USDT", "ETH", "BTC") else ("120.00" if sale else "30.00")
    subtotal = str(D(total) - D(tax))
    if kind == "credit":
        total, subtotal, tax = "-120.00", "-108.00", "-12.00"
    issuer = name if sale else ("测试办公供应商" if lang == "zh" else "Fictional Office Supplies Ltd")
    customer = ("测试企业客户" if lang == "zh" else "Fictional Business Customer Ltd") if sale else name
    invoice_no = f"SIM-{no}-{kind.upper()}-001"
    euro = no == "04"
    total_shown = str(D(total) + 55) if bad_total else total
    if lang == "zh":
        rows = [[issuer], ["销售方名称", issuer], ["购买方名称", customer],
                ["发票号码", invoice_no], ["开票日期", "2026-08-05"], ["Currency", actual_ccy],
                ["项目", DESCRIPTIONS[variant] if sale else "办公运营服务"],
                ["金额", money(subtotal)], ["税额", money(tax)], ["价税合计", money(total_shown)]]
    else:
        rows = [[issuer], ["Bill To", customer], ["Invoice No", invoice_no],
                ["Invoice Date", "2026-08-05"], ["Currency", actual_ccy],
                ["Description", DESCRIPTIONS[variant] if sale else "Office operating services"],
                ["Subtotal", money(subtotal, euro)], ["Tax", money(tax, euro)],
                ["Total Due", money(total_shown, euro)]]
    if kind == "credit":
        rows.insert(1, ["Credit Note" if lang == "en" else "贷项通知单"])
    if ccy in ("USDT", "ETH") and sale:
        # Deliberately generated syntax fixture, never used to send funds.
        rows += [["Payment Method", "On-chain"], ["Ethereum", "0x" + "1" * (39 if bad_wallet else 40)]]
    if ccy == "BTC" and sale:
        rows += [["Bitcoin", "bc1q" + "x" * 38]]
    rows.append(["SYNTHETIC TEST ONLY - NOT A REAL INVOICE"])
    truth = {"invoice_no": invoice_no, "invoice_date": "2026-08-05", "issuer_name": issuer,
             "customer_name": customer, "currency_settlement": actual_ccy,
             "subtotal": subtotal, "sales_tax": tax, "total_due": total_shown}
    return rows, truth


def make_document(path, rows):
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        w = Workbook(); s = w.active; s.title = "Invoice"
        for row in rows:
            s.append(row)
        s.column_dimensions["A"].width = 42; s.column_dimensions["B"].width = 72
        w.save(path); w.close()
    elif suffix == ".docx":
        d = Document()
        for row in rows:
            d.add_paragraph(": ".join(row))
        d.save(path)
    else:
        d = fitz.open(); p = d.new_page(width=720, height=900)
        for i, row in enumerate(rows):
            text = ": ".join(row)
            chinese = any(ord(c) > 127 for c in text)
            p.insert_text((35, 45 + i * 32), text, fontsize=12,
                          fontname="china-s" if chinese else "helv")
        if suffix == ".png":
            p.get_pixmap(matrix=fitz.Matrix(2, 2)).save(path)
        else:
            d.save(path)
        d.close()


def isolate(root):
    paths = {"DATA_ROOT": root, "DATA_DIR": root, "STORAGE_DIR": root / "storage",
             "DB_PATH": root / "app.db", "UPLOAD_DIR": root / "storage/uploads",
             "EXPORT_DIR": root / "exports", "BACKUP_DIR": root / "backups",
             "PAGE_CACHE_DIR": root / "cache/pages", "FX_RATES_PATH": root / "fx.json",
             "FX_UPDATE_STATE_PATH": root / "fx-update-state.json"}
    for key, value in paths.items():
        setattr(config, key, value)
        (value.parent if value.suffix else value).mkdir(parents=True, exist_ok=True)
    config.FX_AUTO_FETCH = False
    db._initialized = False
    db.init_db()


class Scenario:
    def __init__(self, b, root):
        self.b = b; self.root = root / b[0]; self.root.mkdir()
        self.inputs = self.root / "inputs"; self.inputs.mkdir()
        isolate(self.root / "runtime")
        from gateway.main import app
        self.client = TestClient(app, raise_server_exceptions=False)
        self.result = {"id": b[0], "name": b[1], "language": b[2], "sector": b[3],
                       "currency": b[4], "format": b[5], "checks": [], "documents": [],
                       "manual_corrections": [], "requests": [], "complete": False}
        self.account_names = {x["code"]: x["account"] for x in self.call("GET", "/api/ledger/accounts")[1]["accounts"]}

    def check(self, label, ok, actual=None, expected=None, category="workflow"):
        self.result["checks"].append({"name": label, "ok": bool(ok), "actual": actual,
                                      "expected": expected, "category": category})

    def call(self, method, url, body=None, **kwargs):
        r = self.client.request(method, url, json=body, **kwargs)
        try:
            data = r.json()
        except ValueError:
            data = {"bytes": len(r.content), "content_type": r.headers.get("content-type", ""),
                    "error": r.text[:300] if r.status_code >= 400 else None}
        self.result["requests"].append({"method": method, "url": url, "body": body,
                                         "status": r.status_code, "response": data})
        return r.status_code, data

    def post(self, url, body, label):
        status, data = self.call("POST", url, body)
        self.check(label, status == 200 and "error" not in data, data)
        if status != 200 or "error" in data:
            raise RuntimeError(f"{label}: HTTP {status}: {data}")
        return data

    def account(self, code):
        name = self.account_names[code]
        return name if name.startswith(code) else f"{code} {name}"

    def upload(self, path, doc_type="invoice"):
        status, result = self.call("POST", "/api/upload", files={"files": (path.name, path.read_bytes())},
                                   data={"doc_type": doc_type})
        rows = result.get("results", [])
        self.check(f"上传 {path.name}", status == 200 and bool(rows) and all("error" not in x for x in rows), result, category="upload")
        return [x["file_hash_full"] for x in rows if "file_hash_full" in x]

    def extract(self, kind, fmt=None, **kw):
        rows, truth = invoice_rows(self.b, kind, **kw)
        p = self.inputs / f"{kind}.{fmt or self.b[5]}"
        make_document(p, rows)
        hashes = self.upload(p)
        if len(hashes) != 1:
            raise RuntimeError(f"Expected one invoice for {p.name}, got {len(hashes)}")
        h = hashes[0]; inv = db.get_invoice(h)
        extracted = {k: inv.f(k).value for k in truth}
        self.result["documents"].append({"file": str(p.relative_to(self.root)), "hash": h, "truth": truth,
            "extracted": extracted, "parse_status": inv.parse_status, "parse_method": inv.parse_method,
            "ocr_used": inv.ocr_used, "ocr_engine": inv.ocr_engine,
            "issues": [vars(x) for x in inv.issues], "payments": [vars(x) for x in inv.payments],
            "classification": vars(inv.classification)})
        for key, expected in truth.items():
            actual = extracted[key]
            if key in ("subtotal", "sales_tax", "total_due"):
                try:
                    ok = D(str(actual)) == D(expected)
                except Exception:
                    ok = False
            else:
                ok = actual == expected
            self.check(f"{kind} 原始识别 {key}", ok, actual, expected, "extraction")
        raw = self.client.get(f"/api/review/{h}/original")
        self.check(f"{kind} 原件下载字节一致", raw.status_code == 200 and raw.content == p.read_bytes(),
                   {"http_status": raw.status_code, "uploaded_sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
                    "download_sha256": hashlib.sha256(raw.content).hexdigest(), "download_magic": repr(raw.content[:8])})
        return h, truth, p

    def correct(self, h, truth):
        # Explicitly simulate human correction, retain all original failures above.
        for key, expected in truth.items():
            actual = db.get_invoice(h).f(key).value
            if str(actual) != str(expected):
                self.result["manual_corrections"].append({"hash": h, "field": key, "before": actual, "after": expected})
                self.post(f"/api/review/{h}/field", {"field": key, "value": expected, "by": "scenario-reviewer",
                          "reason": "核对独立模拟原件后的人工更正"}, f"人工更正 {key}")

    def approve(self, h, code):
        self.post(f"/api/review/{h}/classify", {"category": "Synthetic reviewed category", "account": self.account(code),
                   "by": "scenario-reviewer", "reason": "模拟审核，科目由测试场景指定"}, "人工确认科目")
        status, data = self.call("POST", f"/api/review/{h}/action", {"action": "Approved", "by": "scenario-reviewer"})
        self.check("正常审核通过", status == 200, data, category="review")
        if status != 200:
            self.result.setdefault("forced_reviews", []).append({"hash": h, "reason": data})
            self.post(f"/api/review/{h}/action", {"action": "Approved", "by": "scenario-reviewer", "force": True,
                      "reason": "仅测试：显式覆盖风险以继续检查后续链路；原阻断已记录"}, "显式覆盖风险继续测试")

    def manual(self, debit, credit, amount, label, activity=None):
        return self.post("/api/ledger/manual", {"date": "2026-08-20", "memo": label,
             "lines": [{"account": self.account(debit), "debit": str(amount)},
                       {"account": self.account(credit), "credit": str(amount)}],
             "activity": activity, "by": "scenario-reviewer"}, label)

    def run(self):
        no, name, lang, sector, ccy, fmt, variant = self.b
        for url in ("/", "/review", "/reconcile", "/ledger", "/reports", "/counterparties"):
            self.check(f"页面 {url}", self.client.get(url).status_code == 200, category="http")
        self.post("/api/counterparties/register", {"name": name, "kind": "self", "note": "虚拟测试主体"}, "登记我方主体")
        self.post("/api/ledger/opening", {"date": "2026-07-31", "other_lines": [
            {"account": self.account("1002"), "amount": "100000", "side": "debit"},
            {"account": self.account("3000"), "amount": "100000", "side": "credit"}]}, "期初建账")
        if ccy not in ("USD", "USDT", "ETH", "BTC"):
            self.post("/api/fx/rate", {"currency": ccy, "date": "2026-08-01", "rate": RATES[ccy]}, "录入合成测试汇率")
        sale, sale_truth, sale_path = self.extract("sale")
        count_before = db.count_invoices()
        self.upload(sale_path)
        self.check("重复上传不增加发票", db.count_invoices() == count_before, db.count_invoices(), count_before)
        code, data = self.call("POST", "/api/ledger/post", {"file_hash": sale, "direction": "AR", "as_of": "2026-08-05"})
        self.check("未审核禁止过账", code == 400, data)
        self.correct(sale, sale_truth)
        rev = db.get_invoice(sale).rev
        self.post(f"/api/review/{sale}/field", {"field": "invoice_no", "value": sale_truth["invoice_no"], "base_rev": rev}, "提交当前版本")
        code, data = self.call("POST", f"/api/review/{sale}/field", {"field": "invoice_no", "value": "STALE", "base_rev": rev})
        self.check("旧版本覆盖被拒绝", code == 409 and db.get_invoice(sale).f("invoice_no").value == sale_truth["invoice_no"], data)
        self.approve(sale, "4000")
        if ccy in ("USDT", "ETH", "BTC"):
            code, data = self.call("POST", "/api/ledger/post", {"file_hash": sale, "direction": "AR", "as_of": "2026-08-05"})
            self.check("加密币无估值汇率禁止入账", code == 400 and "汇率" in str(data), data, category="crypto")
            self.post("/api/fx/rate", {"currency": ccy, "date": "2026-08-01", "rate": RATES[ccy]}, "录入人工合成估值（非市场价）")
        sale_post = self.post("/api/ledger/post", {"file_hash": sale, "direction": "AR", "as_of": "2026-08-05"}, "销售应计入账")
        again_status, again = self.call("POST", "/api/ledger/post", {"file_hash": sale, "direction": "AR", "as_of": "2026-08-05"})
        self.check("重复过账被拒绝并指出原凭证", again_status == 400 and sale_post["entry_no"] in str(again), again)
        from ledger import store
        entry = store.get_entry(sale_post["entry_no"])
        actual_debit = sum(x.debit for x in entry.lines)
        expected_debit = (D(sale_truth["total_due"]) * D(RATES[ccy])).quantize(D("0.01"))
        self.check("销售本位币金额独立核对", actual_debit == expected_debit, actual_debit, expected_debit, "accounting")
        self.result["expected_sale_usd"] = expected_debit
        self.result["actual_sale_usd"] = actual_debit

        purchase, purchase_truth, _ = self.extract("purchase", fmt="pdf" if fmt == "png" else fmt)
        self.correct(purchase, purchase_truth)
        purchase_account = "1405" if variant in ("inventory", "manufacturing", "trading") else "6401" if variant in ("restaurant", "logistics") else "6900"
        self.approve(purchase, purchase_account)
        self.post("/api/ledger/post", {"file_hash": purchase, "direction": "AP", "as_of": "2026-08-05"}, "采购应计入账")
        expected_purchase = (D(purchase_truth["total_due"]) * D(RATES[purchase_truth["currency_settlement"]])).quantize(D("0.01"))

        # Partial settlement tests balance protection; these are test cash receipts,
        # not assertions that crypto wallets are supported as a subledger.
        half = (actual_debit / 2).quantize(D("0.01"))
        self.post("/api/ledger/settle", {"file_hash": sale, "cash_amount": str(half), "settle_amount": str(half), "date": "2026-08-10"}, "分次收款第一笔")
        self.post("/api/ledger/settle", {"file_hash": sale, "cash_amount": str(actual_debit-half), "date": "2026-08-11"}, "分次收款尾款")
        status, data = self.call("POST", "/api/ledger/settle", {"file_hash": sale, "cash_amount": "1", "date": "2026-08-12"})
        self.check("结清后重复结算被拒绝", status == 400, data)
        self.post("/api/ledger/settle", {"file_hash": purchase, "cash_amount": str(expected_purchase), "date": "2026-08-12"}, "支付供应商")

        # Bank statement is an actual uploaded CSV, with independently specified rows.
        p = self.inputs / "bank-statement.csv"
        bank_rows = [["Date", "Description", "Debit", "Credit", "Balance"],
             ["2026-08-10", f"{sale_truth['invoice_no']} customer payment", "", str(half), str(D("100000")+half)],
             ["2026-08-11", f"{sale_truth['invoice_no']} final payment", "", str(actual_debit-half), str(D("100000")+actual_debit)],
             ["2026-08-12", f"{purchase_truth['invoice_no']} supplier payment", str(expected_purchase), "", str(D("100000")+actual_debit-expected_purchase)],
             ["2026-08-13", "Bank fee", "12.50", "", str(D("100000")+actual_debit-expected_purchase-D("12.50"))]]
        if lang == "zh":
            bank_rows[0] = ["日期", "摘要", "支出", "收入", "余额"]
            for row in bank_rows[1:]:
                row[1] = row[1].replace("customer payment", "客户首款").replace("final payment", "客户尾款").replace("supplier payment", "供应商付款").replace("Bank fee", "银行手续费")
        with p.open("w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(bank_rows)
        hs = self.upload(p, "statement")
        if hs:
            st = db.get_invoice(hs[0])
            dump(self.root / "statement.json", {"fields": {k: v.value for k,v in st.fields.items()}, "transactions": [vars(t) for t in st.transactions]})
            self.check("流水识别4笔交易", len(st.transactions) == 4, len(st.transactions), 4, "statement")
            if len(st.transactions) == 4:
                for idx, row in enumerate(bank_rows[1:]):
                    t = st.transactions[idx]
                    ok = D(t.income or 0) == D(row[3] or 0) and D(t.expense or 0) == D(row[2] or 0) and D(t.balance or 0) == D(row[4])
                    self.check(f"流水第{idx+1}笔收支余额", ok, vars(t), row, "statement")
                self.post(f"/api/review/{hs[0]}/field", {"field": "currency_settlement", "value": "USD", "reason": "原始合成流水以USD计价，人工补充币种"}, "补充流水币种")
                self.post("/api/ledger/post-statement", {"stmt_hash": hs[0], "index": 3, "counter_account": self.account("6603"), "activity": "operating"}, "银行手续费流水入账")
            self.call("POST", "/api/reconcile/run", {})
            _, matches = self.call("GET", "/api/reconcile/matches")
            self.result["reconciliation_candidates"] = matches

        extras = []
        if purchase_account == "1405":
            cost = (D(purchase_truth["subtotal"]) * D(RATES[ccy]) / 2).quantize(D("0.01"))
            self.manual("6401", "1405", cost, "售出一半库存结转成本"); extras.append(("cost", str(cost)))
        if variant in ("manufacturing", "health"):
            self.manual("1500", "1002", "2400", "设备购置", "investing")
            self.manual("6602", "1509", "100", "当月折旧"); extras.append(("expense", "100"))
        if variant == "education":
            self.manual("6601", "1002", "150", "培训讲师薪酬", "operating"); extras.append(("expense", "150"))
        if variant == "property":
            self.manual("1002", "2202", "500", "预收服务款不计收入", "operating")
        if variant == "lending":
            self.manual("1002", "2280", "1000", "测试借款本金暂列其它应付款（系统未预设借款科目）", "financing")
            self.manual("6606", "1002", "20", "融资利息支出", "operating"); extras.append(("expense", "20"))
        if variant in ("fund", "broker"):
            # Deliberate generic-accounting boundary: not a fund NAV/client-asset engine.
            self.manual("1002", "2280", "5000", "代收待付本金计负债不计收入", "operating")
            self.manual("2280", "1002", "5000", "代收款原额转出", "operating")
        self.result["extra_business_events"] = extras
        status, data = self.call("POST", "/api/ledger/manual", {"date": "2026-08-20", "lines": [
            {"account": self.account("6900"), "debit": "100"}, {"account": self.account("1002"), "credit": "99"}], "activity": "operating"})
        self.check("借贷不平拒绝入账", status == 400, data)

        # Independent profit oracle, not just the system's internally balanced checks.
        revenue = (D(sale_truth["subtotal"]) * D(RATES[ccy])).quantize(D("0.01"))
        expense = D("0") if purchase_account == "1405" else (D(purchase_truth["subtotal"]) * D(RATES[purchase_truth["currency_settlement"]])).quantize(D("0.01"))
        expected_profit = revenue - expense - D("12.50") - sum(D(x[1]) for x in extras)
        _, rep = self.call("GET", "/api/reports")
        dump(self.root / "reports-before-close.json", rep)
        self.check("关账前五项勾稽", rep.get("can_issue") is True, rep.get("checks"), category="reports")
        self.result["expected_profit_usd"] = expected_profit
        self.result["income_statement"] = rep.get("income_statement")
        actual_profit = rep.get("income_statement", {}).get("net_income")
        self.check("净利润独立预期核对", actual_profit is not None and D(str(actual_profit)) == expected_profit,
                   actual_profit, expected_profit, "accounting")
        self.post("/api/ledger/close", {"period": "2026-08"}, "期末关账")
        _, rep = self.call("GET", "/api/reports")
        dump(self.root / "reports-after-close.json", rep)
        self.check("关账后五项勾稽", rep.get("can_issue") is True, rep.get("checks"), category="reports")
        status, exp = self.call("POST", "/api/reports/export", {})
        self.result["report_exported"] = status == 200 and "download" in exp
        self.check("导出财务报表", self.result["report_exported"], exp, category="reports")
        if self.result["report_exported"]:
            r = self.client.get(exp["download"])
            self.check("报表下载", r.status_code == 200 and r.content[:2] == b"PK", category="reports")
            (self.root / "financial-statements.xlsx").write_bytes(r.content)
            w = load_workbook(io.BytesIO(r.content), data_only=False)
            self.check("Excel七张表完整", len(w.sheetnames) == 7, w.sheetnames, 7, "reports")
            vals = [str(c.value) for row in w["分录明细 凭证轨迹"].iter_rows() for c in row if c.value is not None]
            self.check("Excel可追溯发票号码", any(sale_truth["invoice_no"] in v for v in vals), category="reports")
            w.close()
        else:
            self.check("勾稽失败时阻止报表导出", not rep.get("can_issue") and status == 400, exp, category="safeguard")
        status, data = self.call("POST", "/api/ledger/manual", {"date": "2026-08-25", "lines": [
            {"account": self.account("6900"), "debit": "1"}, {"account": self.account("1002"), "credit": "1"}], "activity": "operating"})
        self.check("已关账期间拒绝新增凭证", status == 400, data)
        self.result["complete"] = True

    def finish(self):
        self.client.close()
        self.result["passed"] = sum(c["ok"] for c in self.result["checks"])
        self.result["failed"] = sum(not c["ok"] for c in self.result["checks"])
        dump(self.root / "result.json", self.result)
        return self.result


def baseline(output):
    import unittest
    isolate(output / "baseline-runtime")
    suite = unittest.defaultTestLoader.discover(str(ROOT / "tests"), pattern="test_*.py")
    with (output / "baseline.log").open("w", encoding="utf-8") as f:
        r = unittest.TextTestRunner(stream=f, verbosity=2).run(suite)
    summary = {"run": r.testsRun, "failures": len(r.failures), "errors": len(r.errors), "skipped": len(r.skipped),
               "failure_details": [(str(t), s) for t,s in r.failures+r.errors], "skips": [(str(t), s) for t,s in r.skipped]}
    dump(output / "baseline.json", summary)
    print(json.dumps(summary, ensure_ascii=False), flush=True)
    return r.wasSuccessful()


def boundary_probes(output):
    """Additional original-file probes, independent from the 16 business ledgers."""
    b = ("90", "Fictional Boundary Laboratory Ltd", "en", "独立边界验证", "USD", "pdf", "subscription")
    s = Scenario(b, output)

    def original(label, rows, truth, fmt="pdf"):
        # Change invoice IDs so that unrelated probes do not trigger deduplication.
        invoice_no = "PROBE-" + label.upper()
        truth = dict(truth, invoice_no=invoice_no)
        rows = [([row[0], invoice_no] if len(row) == 2 and row[0] in ("Invoice No", "发票号码") else row) for row in rows]
        p = s.inputs / f"{label}.{fmt}"; make_document(p, rows)
        hs = s.upload(p)
        if not hs:
            raise RuntimeError(f"{label}: upload failed")
        h = hs[0]; inv = db.get_invoice(h)
        s.result["documents"].append({"file": str(p.relative_to(s.root)), "hash": h, "truth": truth,
             "extracted": {k: inv.f(k).value for k in truth}, "issues": [vars(x) for x in inv.issues],
             "payments": [vars(x) for x in inv.payments]})
        return h, truth, p

    def guard(label, fn):
        try:
            fn()
        except Exception:
            s.check(label + "执行完整", False, traceback.format_exc(), category="probe-error")

    def credit():
        rows, truth = invoice_rows(b, "credit")
        h, truth, _ = original("credit", rows, truth)
        inv = db.get_invoice(h)
        s.check("贷项负总额保留", inv.f("total_due").value == D("-120"), inv.f("total_due").value, "-120", "extraction")
        s.correct(h, truth); s.approve(h, "4000")
        status, r = s.call("POST", "/api/ledger/post", {"file_hash": h, "direction": "AR"})
        s.check("销售贷项入账", status == 200, r, category="accounting")
        if status != 200:
            s.result.setdefault("blocked_checks", []).append("负数贷项未入账，无法验证收入与销项税反向分录")
            return
        from ledger import store
        e = store.get_entry(r["entry_no"])
        by_code = {line.account.split()[0]: (line.debit, line.credit) for line in e.lines}
        s.check("贷项收入与销项税同时冲减", by_code.get("4000") == (D("108"), D("0")) and by_code.get("2210") == (D("12"), D("0")), by_code,
                  {"4000": ["108", "0"], "2210": ["12", "0"]}, "accounting")

    def tax_error():
        rows, truth = invoice_rows(b, bad_total=True)
        h, _, _ = original("bad-total", rows, truth)
        issues = [x.code for x in db.get_invoice(h).issues]
        s.check("净额加税不等总额有提示", any("MISMATCH" in x for x in issues), issues, category="safeguard")

    def wallet_error():
        cb = list(b); cb[4] = "USDT"
        rows, truth = invoice_rows(tuple(cb), bad_wallet=True)
        h, _, _ = original("bad-wallet", rows, truth)
        inv = db.get_invoice(h)
        issues = [x.code for x in inv.issues]
        s.check("39位EVM地址有风险提示", "WALLET_FORMAT" in issues, issues, category="crypto")
        s.check("错误钱包未标有效", any(not p.valid_address for p in inv.payments), [vars(p) for p in inv.payments], category="crypto")

    def missing_currency():
        rows, truth = invoice_rows(b)
        rows = [r for r in rows if r[0] != "Currency"]
        h, _, _ = original("missing-currency", rows, truth)
        before = db.get_invoice(h)
        status, data = s.call("POST", f"/api/review/{h}/action", {"action": "Approved"})
        s.check("没有币种不能无提示审核", status == 400 or any("CURRENCY" in x.code for x in before.issues),
                  {"status": status, "response": data, "inferred_currency": before.f("currency_settlement").value,
                   "issues": [vars(x) for x in before.issues]}, category="safeguard")

    def multi_pdf():
        paths = []; expected = []
        for i in range(2):
            rows, truth = invoice_rows(b)
            no = f"MULTI-PROBE-{i}"
            rows = [([r[0], no] if r[0] == "Invoice No" else r) for r in rows]
            p = s.inputs / f"multi-source-{i}.pdf"; make_document(p, rows); paths.append(p); expected.append(no)
        d = fitz.open()
        for p in paths:
            with fitz.open(p) as part:
                d.insert_pdf(part)
        p = s.inputs / "two-invoices.pdf"; d.save(p); d.close()
        hs = s.upload(p)
        actual = sorted(str(db.get_invoice(h).f("invoice_no").value) for h in hs)
        s.check("两页两发票自动拆分", actual == expected, actual, expected, "extraction")
        if hs:
            inv = db.get_invoice(hs[0])
            r = s.client.get(f"/api/collection/{inv.source_file_hash}/original")
            s.check("合集原始文件可下载", r.status_code == 200 and r.content == p.read_bytes())

    def fx_settlement():
        fb = list(b); fb[4] = "EUR"
        rows, truth = invoice_rows(tuple(fb))
        h, truth, _ = original("fx-settlement", rows, truth)
        s.correct(h, truth); s.approve(h, "4000")
        for date, rate in (("2026-08-05", "1.10"), ("2026-08-15", "1.20")):
            s.post("/api/fx/rate", {"currency": "EUR", "date": date, "rate": rate}, "外币合成汇率")
        s.post("/api/ledger/post", {"file_hash": h, "direction": "AR", "as_of": "2026-08-05"}, "外币应收入账")
        result = s.post("/api/ledger/settle", {"file_hash": h, "cash_amount": "1200", "cash_currency": "EUR",
                        "settle_amount": "1320", "date": "2026-08-15", "diff_reason": "fx_gain_loss"}, "外币实际币种结算")
        from ledger import store
        e = store.get_entry(result["entry_no"])
        gain = sum(l.credit-l.debit for l in e.lines if l.account.startswith("6607 "))
        s.check("汇率变化产生120美元收益", gain == D("120"), gain, "120", "accounting")

    def crypto_settlement():
        cb = list(b); cb[4] = "ETH"
        rows, truth = invoice_rows(tuple(cb))
        h, truth, _ = original("eth-settlement", rows, truth)
        s.correct(h, truth); s.approve(h, "4000")
        s.post("/api/fx/rate", {"currency": "ETH", "date": "2026-08-01", "rate": "2400"}, "ETH人工合成估值")
        s.post("/api/ledger/post", {"file_hash": h, "direction": "AR", "as_of": "2026-08-05"}, "ETH计价服务入账")
        r = s.post("/api/ledger/settle", {"file_hash": h, "cash_amount": "1.23456789", "cash_currency": "ETH",
                    "settle_amount": "2962.96", "date": "2026-08-15"}, "ETH原币高精度结算")
        from ledger import store
        e = store.get_entry(r["entry_no"])
        value = sum(l.debit for l in e.lines if l.account.startswith("1002 "))
        s.check("ETH原币乘估值再舍入", value == D("2962.96"), value, "2962.96", "crypto")
        s.result["crypto_cash_account_observation"] = {"lines": [vars(l) for l in e.lines], "memo": e.memo,
            "limitation": "当前HTTP结算默认记银行科目；数值换算通过不等于存在钱包/代币数量台账或链上对账。"}

    def reconcile_flow():
        rows, truth = invoice_rows(b)
        h, truth, _ = original("reconcile", rows, truth)
        s.correct(h, truth); s.approve(h, "4000")
        s.post("/api/ledger/post", {"file_hash": h, "direction": "AR"}, "对账用销售入账")
        path = s.inputs / "reconcile-bank.csv"
        with path.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows([["Date", "Description", "Debit", "Credit", "Balance"],
                ["2026-08-15", "Fictional Business Customer Ltd " + truth["invoice_no"], "", "1200.00", "1200.00"]])
        hs = s.upload(path, "statement")
        s.post(f"/api/review/{hs[0]}/field", {"field": "currency_settlement", "value": "USD"}, "对账流水币种")
        s.post("/api/reconcile/run", {}, "运行对账匹配")
        _, data = s.call("GET", "/api/reconcile/matches")
        dump(s.root / "reconcile-candidates.json", data)
        # Select by serialized source hash rather than by candidate order.
        candidates = [m for m in data.get("matches", []) if h in json.dumps(m)]
        s.check("已识别发票与流水存在匹配候选", bool(candidates), data, category="reconciliation")
        s.result["reconcile_invoice_hash"] = h
        s.result["reconcile_statement_hash"] = hs[0]
        if candidates:
            m = candidates[0]; mid = m.get("id", m.get("match_id"))
            s.post("/api/reconcile/confirm", {"match_id": mid}, "人工确认匹配")
            s.post("/api/ledger/settle-from-match", {"file_hash": h}, "从已确认匹配结算")
            code, data = s.call("POST", "/api/ledger/settle-from-match", {"file_hash": h})
            s.check("重复对账结算被拒绝", code == 400, data, category="safeguard")

    for label, fn in [("贷项", credit), ("税额异常", tax_error), ("钱包", wallet_error),
                      ("缺币种", missing_currency), ("合集", multi_pdf), ("外币", fx_settlement),
                      ("ETH", crypto_settlement), ("对账", reconcile_flow)]:
        guard(label, fn)
    # Bad inputs use the same upload handler as the UI, including mixed-batch isolation.
    status, data = s.call("POST", "/api/upload", files=[("files", ("empty.pdf", b"")),
        ("files", ("unsupported.exe", b"test")), ("files", ("broken.pdf", b"not a pdf"))])
    s.check("无效上传返回错误或待人工失败记录", status == 200 and len(data.get("results", [])) == 3 and
            all("error" in r or r.get("parse_failed") for r in data.get("results", [])), data, category="safeguard")
    from core import maintenance
    import sqlite3
    backup = maintenance.backup_db(config.DB_PATH, config.BACKUP_DIR)
    s.check("隔离账套备份完整性", backup is not None and maintenance.integrity_check(backup), str(backup), category="backup")
    if backup:
        with sqlite3.connect(f"file:{backup}?mode=ro", uri=True) as conn:
            count = conn.execute("SELECT count(*) FROM invoices").fetchone()[0]
        s.check("备份与原测试账套单据数量相同", count == db.count_invoices(), count, db.count_invoices(), "backup")
    s.result["complete"] = True
    result = s.finish(); dump(output / "boundary.json", result)
    print(json.dumps({"passed": result["passed"], "failed": result["failed"]}), flush=True)
    return result["failed"] == 0


def main():
    p = argparse.ArgumentParser(); p.add_argument("--output", required=True, type=Path)
    p.add_argument("--ids", default=""); p.add_argument("--baseline", action="store_true")
    p.add_argument("--probes", action="store_true")
    args = p.parse_args(); args.output.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.ERROR)
    if args.baseline:
        raise SystemExit(0 if baseline(args.output) else 1)
    if args.probes:
        raise SystemExit(0 if boundary_probes(args.output) else 1)
    dump(args.output / "manifest.json", {"synthetic_only": True, "rates_are_not_market_data": True,
        "businesses": BUSINESSES, "rates": RATES, "transport": "real FastAPI routes via in-process ASGI TestClient"})
    results = []
    for b in BUSINESSES:
        if args.ids and b[0] not in args.ids.split(","):
            continue
        print(f"START {b[0]} {b[1]}", flush=True)
        s = None
        try:
            s = Scenario(b, args.output); s.run()
        except Exception:
            detail = traceback.format_exc()
            if s:
                s.result["fatal"] = detail
            else:
                results.append({"id": b[0], "name": b[1], "complete": False,
                    "passed": 0, "failed": 1, "checks": [], "fatal": detail})
            print(detail, flush=True)
        finally:
            if s:
                results.append(s.finish())
                print(f"END {b[0]} passed={s.result['passed']} failed={s.result['failed']} complete={s.result['complete']}", flush=True)
        dump(args.output / "summary.json", results)
    raise SystemExit(1 if any(r["failed"] or not r["complete"] for r in results) else 0)


if __name__ == "__main__":
    main()
