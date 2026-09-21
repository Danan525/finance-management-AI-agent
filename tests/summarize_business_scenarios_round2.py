"""Produce round-two evidence without hiding extraction failures or manual edits."""
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]


def main(out):
    read = lambda p: json.loads((out / p).read_text())
    cases = read("final/summary.json")
    previous = read("previous-cohort/summary.json")
    baseline = read("baseline/baseline.json")
    assert len(cases) >= 20 and all(s["complete"] for s in cases)
    repeats = [s for s in cases if s["previous_problem_targets"]]
    assert len(repeats) >= 5
    checks = [(s["id"], c) for s in cases for c in s["checks"]]
    extraction = [(i, c) for i, c in checks if c["category"] == "extraction"]
    failed = [(i, c) for i, c in checks if not c["ok"]]
    fixes = [
        ["F01", "中文OCR金额缩小", "已修复并扩展", "仅对完整数字token规范千位分隔及全角标点，保留原始OCR文本；有效欧式格式和币种小数精度不变。", "23、42"],
        ["F02", "贷项负号及入账", "已修复", "保留负号；应收/应付贷项反向净额与税额，支持分次退款、超额拒绝、退款红冲及控制科目核对。", "37、38；专项回归"],
        ["F03", "中文购买方标签串入名称", "已修复", "优先完整匹配购买方名称等标签；覆盖PDF、图片、Word、Excel。", "22、23、25、27、38、42"],
        ["F04", "中文金额漏作未税额", "已修复指定歧义边界", "仅在候选唯一且与明确税额、总额相加一致时识别为未税额；冲突或多候选保持人工复核。", "22、23、25、27、38、42"],
        ["F05", "USDT币种遗漏", "已修复", "支持USDT/USDC完整币种，不截成USD；USDT按合成估值0.997核对。", "34；专项USDC回归"],
        ["F06", "英文OCR名称空格丢失", "部分改善，仍有边界", "仅以唯一已建档全名恢复空格，保留原文；两处未建档客户名称仍需人工校正。", "旧08、新28/44；专项唯一/歧义匹配"],
        ["F07", "Word原件下载成PDF", "已修复", "原件接口优先返回上传源文件并校验存储路径；预览仍用转换件。", "27；专项越界路径回归"],
        ["F08", "借款筹资现金流E4阻断", "已修复场景及映射", "增加2001短期借款/2501长期借款；间接法剔除明确投资/筹资分录中的营运资金变动，兼容原2280场景。", "32；专项三种科目及还本回归"],
        ["C01", "代币收付被记作银行", "增加已知币种拦截", "结算入口拒绝BTC/ETH/USDT/USDC/DAI/SOL/WBTC/WETH原币现金；真实USD收款仍可结算。未实现链上资产子账。", "34、35、36；专项回归"],
    ]
    cols = ["编号", "虚拟主体", "行业", "语言", "币种", "主件", "旧问题复测", "通过/未达预期", "人工改字段", "报表"]
    rows = [[s["id"], s["name"], s["sector"], s["language"], s["currency"], s["format"],
             ",".join(s["previous_problem_targets"]) or "—", f'{s["passed"]}/{s["failed"]}',
             len(s["manual_corrections"]), "已导出" if s["report_exported"] else "未导出"] for s in cases]
    table = ["|" + "|".join(cols) + "|", "|" + "---|" * len(cols)]
    for row in rows:
        linked = row.copy(); linked[1] = f'[{row[1]}](final/{row[0]}/result.json)'
        table.append("|" + "|".join(map(str, linked)) + "|")
    fix_table = ["|编号|问题|状态|改动与限制|验证主体|", "|---|---|---|---|---|"]
    fix_table += ["|" + "|".join(row) + "|" for row in fixes]
    files = sorted((out / "final").glob("*/inputs/*"))
    formats = Counter(p.suffix for p in files)
    langs = Counter(s["language"] for s in cases)
    currency = Counter(s["currency"] for s in cases)
    old_fail = sum(s["failed"] for s in previous)
    old_pass = sum(s["passed"] for s in previous)
    changes = sum(len(s["manual_corrections"]) for s in cases)
    report = f"""# 第二轮：24个中英文主体修复与实测

测试日期：2026-09-18。**{len(cases)}个新主体全部完成账务流程并导出报表，其中{len(repeats)}个复测上一轮相似问题。仍有{len(failed)}处未建档英文客户名OCR空格偏差，经过显式人工更正后继续；不宣称全自动识别通过。**

## 实际结果

- 新主体检查{len(checks)}项：{len(checks)-len(failed)}通过、{len(failed)}未达预期。首次提取{len(extraction)}个字段：{sum(c['ok'] for _,c in extraction)}正确、{sum(not c['ok'] for _,c in extraction)}不符；共{changes}处人工更正，原始失败未覆盖。
- 全量回归{baseline['run']}项：{baseline['failures']}失败、{baseline['errors']}错误、{baseline['skipped']}跳过（包含11项新增专项）；全角标点增量后另复查45项OCR/专项用例，全部通过，数量不重复相加。
- 另把上一轮16主体在新隔离目录重跑：{old_pass}通过、{old_fail}未达预期，{sum(bool(s['report_exported']) for s in previous)}份报表导出，{sum(len(s['manual_corrections']) for s in previous)}处字段校正。上一轮独立边界中“ETH原币可结算”的旧预期已被本轮拒绝测试替代，未机械重用旧预期。
- 新主体{langs['zh']}个中文、{langs['en']}个英文；币种{', '.join(f'{c}×{n}' for c,n in currency.items())}。
- 保留{len(files)}份合成输入：{', '.join(f'{k}×{v}' for k,v in sorted(formats.items()))}；24份最终Excel报表，每份7张工作表。
- 真实接口完成上传→首次提取→复核/审核→入账→结算/退款→关账→勾稽→报表导出。销售折算和最终净利润与独立Decimal预期核对，不能只用系统自己的借贷平衡证明正确。

## 主体清单

{chr(10).join(table)}

行业含软件订阅、零售电商、餐饮、制造、物流、设计、教育、医疗、物业、跨境贸易、通信、酒店、新能源工程、文化传媒、非营利培训，以及资管、融资、经纪、支付、USDT/ETH/BTC计价服务。
这些是普通财务业务场景的模拟，不表示覆盖各行业所有专门业务或税务规则。

## 上轮问题与修复证据

{chr(10).join(fix_table)}

相似问题主体共10个：22、23、25、27、28、32、34、37、38、42。既更换主体名称，也更换法币金额、部分行业和单据格式；没有直接复制上一轮结果。

## 覆盖的流程和边界

- 中英文PDF/图片OCR/Word/Excel发票，中文/英文CSV流水；欧式金额、全角数字标点、负数、八位币种数量、币种缺失与合成汇率换算。
- 主体登记、期初建账、源文件下载字节核对、重复上传、审核前禁止入账、版本冲突、重复过账/结算拦截、两次收款、采购付款、控制科目与往来明细核对。
- 库存及成本、设备购置及折旧、薪酬、预收款、代收本金负债及转出、借款及利息；独立核对最终净利润、关账后五项勾稽和7张报表工作表。
- 37/38分别完成销售/采购贷项：原文件负数提取、净额与税额反向、分次退款、退款现金方向、结清后的往来控制、重新关账出表。专项还验证超额退款拒绝、退款红冲后余额恢复、不可抵扣税额并成本。
- 34/35/36分别验证USDT/ETH/BTC：第一份模拟合同实际收到USD，第二份代币计价应收保持未结；尝试代币原币结算时HTTP400且凭证数量不变。
- 专项验证原件下载路径边界、注册名称唯一匹配及歧义不自动改名、重复/冲突“金额”不武断推成未税额、2001/2501/2280借款与还本勾稽。

## 剩余问题与适用范围

1. **两处首次提取失败仍保留**：28/44的未建档客户`Fictional Business Customer Ltd`被识别成`Fictional BusinessCustomerLtd`；经真实审核接口各更正一次。已有唯一主档的名称可辅助恢复；这不是独立OCR模型准确率提升，也不根据驼峰猜测名称边界。
2. “金额”含义不明确或有多候选时继续要求人工核对；不保证任意扫描质量/任意版式均能自动识别。初始分类中Unclassified项通过审核显式选科目，本轮不把它算成自动分类通过。
3. 贷项覆盖独立贷项过账与现金退款，**未实现贷项自动抵扣其他发票**；贷项业务与普通凭证红冲仍是不同动作。
4. 数字资产拦截仅针对配置列出的已知代码及发票结算入口；不能宣称所有币种、手工凭证和流水入口都有数字资产识别。无链上钱包数量账、交易哈希核验、Gas归集、质押奖励、基金NAV、证券持仓或衍生品估值支持。
5. 本次为进程内ASGI真实接口/业务实测；页面检查可访问，修改页面JS通过语法检查，**未进行浏览器逐按钮和布局测试**。不包含多人同步、权限改造、压力测试或生产部署。
6. 全部虚构公司、票据、金额、汇率和税额，逐主体隔离SQLite、附件与输出，自动汇率关闭并阻止外部TCP连接。没有向真实账套写入测试交易。

## 证据与复跑

- [逐项测试Excel](测试明细.xlsx)：主体、检查、修复、首次提取含来源、人工更正、汇率。
- [新主体完整结果](final/summary.json)、[旧主体重跑](previous-cohort/summary.json)、[全量回归日志](baseline/baseline.log)、[增量OCR回归日志](verification/ocr-regressions.log)、[证据清单及源码哈希](evidence-manifest.json)。
- `final/编号/inputs/`保留源文件，`result.json`保留首次识别、校正、请求响应；`reports-after-close.json`及`financial-statements.xlsx`为最终结果。运行数据库在各自runtime中，本地保留但不放入压缩包。
- [合成样例与结果包](合成样例与测试结果.zip)仅打包虚拟测试证据和复跑脚本，不含实际运行配置或真实业务数据库。`cohort/`保留本轮最初发现3处偏差的过程证据，最终统计取`final/`。

在项目根目录使用全新的输出目录：

```bash
.venv/bin/python tests/run_business_scenarios_round2.py --output data/round2-retest/final
.venv/bin/python tests/run_business_scenarios_round2.py --output data/round2-retest/baseline --baseline
.venv/bin/python tests/run_business_scenarios.py --output data/round2-retest/previous-cohort
.venv/bin/python tests/summarize_business_scenarios_round2.py data/round2-retest
```

脚本保留全部未达预期项，因此当前24主体测试退出码为1（两处OCR名称空格偏差），不能把这个退出码改成成功来隐藏剩余问题。可用`--ids 37,38`单独验证贷项退款。
"""
    (out / "测试报告.md").write_text(report, encoding="utf-8")
    sheets = {
        "主体总览": [cols] + rows,
        "逐项检查": [["主体", "类型", "检查", "结果", "实际", "期望"]] +
            [[i, c["category"], c["name"], "通过" if c["ok"] else "未达预期", json.dumps(c["actual"], ensure_ascii=False),
              json.dumps(c["expected"], ensure_ascii=False)] for i,c in checks],
        "修复和边界": [["编号", "问题", "状态", "改动与限制", "证据"]] + fixes,
        "首次提取": [["主体", "文件", "字段", "期望", "首次结果", "方法", "原始字段", "处理来源", "注释"]] +
            [[s["id"],d["file"],k,str(v),str(d["extracted"].get(k)),d["parse_method"],
              str(d["field_provenance"][k]["raw"]),d["field_provenance"][k]["source"],d["field_provenance"][k]["note"]]
             for s in cases for d in s["documents"] for k,v in d["truth"].items()],
        "人工更正": [["主体", "字段", "更正前", "更正后"]] +
            [[s["id"],c["field"],c["before"],c["after"]] for s in cases for c in s["manual_corrections"]],
        "合成汇率": [["币种", "测试USD估值（非市场价）"]] + list(read("final/manifest.json")["rates"].items()),
    }
    w = Workbook(); w.remove(w.active)
    for title, content in sheets.items():
        sh = w.create_sheet(title)
        for row in content:
            sh.append(row)
        sh.freeze_panes = "A2"; sh.auto_filter.ref = sh.dimensions
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="17365D")
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(64, max(15, max(len(str(c.value or "")) for c in col[:80])+2))
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    w.save(out / "测试明细.xlsx"); w.close()
    sources = ["core/config.py", "extraction/extract/ocr.py", "extraction/parse/generic.py", "extraction/parse/template_rules.py",
               "extraction/pipeline.py", "gateway/main.py", "ledger/accounts.py", "ledger/posting.py", "ledger/settlement.py",
               "ledger/service.py", "ledger/store.py", "reports/mapping.py", "reports/service.py", "web/ledger.html",
               "tests/run_business_scenarios.py", "tests/run_business_scenarios_round2.py", "tests/test_round2_regressions.py",
               "tests/summarize_business_scenarios_round2.py"]
    sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    manifest = {"generated_at": datetime.now(timezone.utc).isoformat(),
                "base_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
                "code_state": "含未提交修复，以source_hashes为准", "source_hashes": {p:sha(ROOT/p) for p in sources},
                "subjects":len(cases), "repeat_subjects":len(repeats), "checks":len(checks), "failed":len(failed),
                "extraction_checks":len(extraction), "manual_corrections":changes, "baseline":baseline,
                "input_hashes":{str(p.relative_to(out)):sha(p) for p in files}}
    (out / "evidence-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    with zipfile.ZipFile(out / "合成样例与测试结果.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for name in ("测试报告.md", "测试明细.xlsx", "evidence-manifest.json", "baseline/baseline.json", "baseline/baseline.log"):
            z.write(out / name, name)
        for folder in ("final", "previous-cohort", "cohort", "verification"):
            for p in sorted((out / folder).rglob("*")):
                rel = p.relative_to(out)
                if p.is_file() and "runtime" not in rel.parts:
                    z.write(p, str(rel))
        for p in sources:
            if p.startswith("tests/"):
                z.write(ROOT / p, "复跑脚本/" + Path(p).name)
    print(json.dumps({k:manifest[k] for k in ("subjects", "repeat_subjects", "checks", "failed", "extraction_checks", "manual_corrections")}, ensure_ascii=False))


if __name__ == "__main__":
    main(Path(sys.argv[1]))
