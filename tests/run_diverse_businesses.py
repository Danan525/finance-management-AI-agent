"""30 distinct businesses: 20 layout families and different accounting paths.

No invoice_rows/make_document reuse: this cohort deliberately varies geometry,
tables, pagination and wording. HTTP harness reuse only; all data is synthetic.
"""
from __future__ import annotations
import argparse
import io
import json
import logging
from pathlib import Path
from decimal import Decimal as D
import traceback

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageFilter
import run_business_scenarios as base
from core import db, config
from ledger import accounts as A

PROFILES = [
 ('Horizon SaaS Partners Ltd','软件订阅/未收尾款','service',0),
 ('湖州新衡精密制造有限公司','制造/设备赊购未付款','asset_unpaid',1),
 ('Orchid Dental Equipment LLC','医疗/设备分期付款','asset_partial',2),
 ('泉州锦帆物流有限公司','物流/赊购车辆后付款','asset_paid',3),
 ('Beacon Cinema Operations Ltd','影院/折旧红冲重提','depreciation',4),
 ('常州青禾机械有限公司','工程机械/处置再撤销','disposal',5),
 ('Ridgeway Research Systems Ltd','研发/软件资产摊销撤销','intangible',6),
 ('嘉兴青麦食品有限公司','食品/多品项库存及成本','inventory',7),
 ('Cobalt Workspace Services Ltd','共享办公/预付租金摊销','prepaid',8),
 ('福州梧桐酒店有限公司','酒店/预收及收入确认','advance',9),
 ('Northwind Staffing Services Ltd','人力服务/薪酬计提分次支付','payroll',10),
 ('绍兴川谷商贸有限公司','商贸/进销项税额结转缴纳','vat',11),
 ('Lantern Leasing Services Ltd','融资服务/借款部分归还','loan',12),
 ('扬州启辰工业有限公司','制造/应付款转长期借款','novation',13),
 ('Fairhaven Holding Services Ltd','控股服务/宣告并支付股利','dividend',14),
 ('南通星芒技术有限公司','科技/非现金股权薪酬','equity_pay',15),
 ('Silver Fern Export GmbH','出口/EUR跨汇率收款','fx',16),
 ('益阳清舟零售有限公司','零售/销售贷项部分退款','credit_ar',17),
 ('Marigold Consulting Ltd','咨询/采购贷项退款','credit_ap',18),
 ('宜昌远岭数字有限公司','数字服务/无形资产退购','asset_credit',19),
 ('Pelican Marketplace Services Ltd','平台/扣手续费净额到账','fee',2),
 ('盐城禾景专业服务有限公司','专业服务/收款预扣税','withholding',11),
 ('Clearbrook Asset Services Ltd','资产服务/赊销设备再收款','credit_disposal',5),
 ('松原青石数字托管有限公司','加密/BTC计价及USD实际收款','btc',9),
 ('Prairie Stablecoin Infrastructure Ltd','加密/USDT计价未收款','usdt',7),
 ('宁德云帆验证技术有限公司','加密/ETH计价未收款','eth',13),
 ('Brookhaven Community Foundation','公益/现金捐赠及项目支出','nonprofit',12),
 ('湘潭白鹭旅行服务有限公司','旅游/银行现金内部转账','transfer',15),
 ('Arcadia Telecom Services Ltd','通信/跨月关账及重开','periods',6),
 ('威海麦湾支付服务有限公司','支付服务/代收本金转出','custody',18),
]
LAYOUTS = [
 '双栏抬头+商品表+右下合计PDF','中文标签在上/值在下PDF','横向纸张双栏地址PDF',
 '宽表中文表头/数据两行PDF','窄幅热敏小票PNG','两页PDF正文续页合计',
 'DOCX双方信息表+商品表','欧式数字与右侧合计PDF','XLSX封面及第二表账单',
 'XLSX合并抬头/分散单元格','DOCX段落与独立汇总表','分列税项CGST/SGST PDF',
 '合计在顶部、项目明细在下PDF','长描述换行与页脚干扰PDF','折扣明细后净额PDF',
 '轻微倾斜扫描PNG','标签值分栏/币种在页头PDF','括号负数贷项PDF',
 'XLSX商品数量单价表及汇总','DOCX贷项通知单及表格',
]


def fixture(path, layout, name, issuer, customer, truth):
    """Make structurally varied documents from truth (never from parser output)."""
    zh = any(ord(c)>127 for c in name)
    net,tax,total = [D(truth[k]) for k in ('subtotal','sales_tax','total_due')]
    def money(v):
        s = f'{v:,.2f}' if v.as_tuple().exponent >= -2 else str(v)
        if layout == 7 or truth['currency_settlement']=='EUR':
            s = s.replace(',', '@').replace('.', ',').replace('@', '.')
        if layout == 17 and v < 0:
            s = '('+s.lstrip('-')+')'
        return s
    no=truth['invoice_no']; ccy=truth['currency_settlement']
    identity = [('销售方名称' if zh else 'Seller',issuer), ('购买方名称' if zh else 'Bill To',customer),
                ('发票号码' if zh else 'Invoice No',no),('开票日期' if zh else 'Invoice Date','2026-08-05'),('Currency',ccy)]
    totals = [('未税金额' if zh else 'Subtotal',money(net)),('税额' if zh else 'Tax',money(tax)),
              ('价税合计' if zh else 'Total Due',money(total))]
    title = ('贷项通知单' if zh else 'Credit Note') if total<0 else ('服务发票' if zh else 'INVOICE')
    if path.suffix == '.xlsx':
        w=Workbook(); s=w.active
        if layout==8:
            s.title='Readme'; s.append(['SYNTHETIC TEST ONLY']); s=w.create_sheet('账单 Invoice'); start=8
        else: s.title='Invoice'; start=3
        s.merge_cells('A1:F1'); s['A1']=title
        for i,(k,v) in enumerate(identity,start): s.cell(i,1,k); s.cell(i,4,v)
        i=start+7
        if layout==18:
            for j,v in enumerate(['Description','Qty','Unit Price','Amount'],1):s.cell(i,j,v)
            for off,q in ((1,D('.4')),(2,D('.6'))):
                for j,v in enumerate(['Synthetic item '+str(off),1,money(net*q),money(net*q)],1):s.cell(i+off,j,v)
            i+=5
        for off,(k,v) in enumerate(totals): s.cell(i+off,3,k); s.cell(i+off,6,v)
        w.save(path); w.close(); return
    if path.suffix=='.docx':
        d=Document(); d.add_heading(title,0)
        if layout==10:
            for k,v in identity:d.add_paragraph(k+': '+v)
        else:
            t=d.add_table(rows=0,cols=2)
            for k,v in identity: a,b=t.add_row().cells; a.text=k;b.text=v
        d.add_paragraph('Synthetic goods and services; not a real invoice.')
        t=d.add_table(rows=0,cols=2)
        for k,v in totals:a,b=t.add_row().cells;a.text=k;b.text=v
        d.save(path);return
    doc=fitz.open(); width=900 if layout==2 else 330 if layout==4 else 720
    page=doc.new_page(width=width,height=900)
    def put(x,y,text,size=12):
        font='china-s' if any(ord(c)>127 for c in str(text)) else 'helv'
        page.insert_text((x,y),str(text),fontname=font,fontsize=size)
    def pair(y,k,v,x=32,gap=270):put(x,y,k+':');put(x+gap,y,v)
    put(28,30,title,17)
    if layout==1:
        for i,(k,v) in enumerate(identity):put(35,65+i*64,k);put(35,91+i*64,v)
        for i,(k,v) in enumerate(totals):pair(450+i*38,k,v)
    elif layout==3:
        for i,(k,v) in enumerate(identity):pair(70+i*32,k,v)
        for i,(k,v) in enumerate(totals):put(35+i*225,330,k);put(35+i*225,365,v)
    elif layout==4:
        for i,(k,v) in enumerate(identity):put(16,70+i*52,k+':',10);put(16,91+i*52,v,10)
        for i,(k,v) in enumerate(totals):pair(395+i*38,k,v,16,170)
    else:
        if layout in (0,2):
            put(30,65,identity[0][0]+':');put(30,90,issuer)
            put(width/2,65,identity[1][0]+':');put(width/2,90,customer)
            for i,(k,v) in enumerate(identity[2:]):pair(145+i*35,k,v)
        elif layout==16:
            put(510,50,'Currency: '+ccy)
            for i,(k,v) in enumerate(identity[:-1]):pair(110+i*46,k,v,35,330)
        else:
            for i,(k,v) in enumerate(identity):pair(65+i*32,k,v)
        if layout==12:
            for i,(k,v) in enumerate(totals):pair(250+i*35,k,v)
            put(30,500,'Item / Quantity / Unit Price / Amount')
            put(30,535,'Consulting services / 1 / '+money(net)+' / '+money(net))
        elif layout==5:
            put(30,310,'Project milestones and delivery details - continued on page 2')
            for i in range(10):put(30,350+i*35,'Work package '+str(i+1)+' completed and accepted.')
            page=doc.new_page(width=720,height=900)
            put(30,40,'Continuation / Page 2')
            for i,(k,v) in enumerate(totals):pair(400+i*40,k,v)
        else:
            page.draw_line((30,280),(width-30,280))
            put(30,305,'Description / Qty / Unit price / Amount')
            put(30,335,'Service A / 1 / '+money(net*D('.4'))+' / '+money(net*D('.4')))
            put(30,365,'Service B / 1 / '+money(net*D('.6'))+' / '+money(net*D('.6')))
            if layout==13:
                for i in range(3):put(30,390+i*20,'Extended scope and acceptance requirements; description continues.',10)
            if layout==14:
                put(30,400,'List price: '+money(net+100));put(30,425,'Commercial discount: 100.00')
            for i,(k,v) in enumerate(totals):
                if layout==11 and i==1:
                    pair(545,'CGST',money(tax/2));pair(580,'SGST',money(tax/2))
                else:pair(500+i*(65 if layout==11 else 40),k,v,300 if layout==0 else 35,200 if layout==0 else 330)
            if layout==13:put(35,720,'Help desk 2026 / Reference 800100 / Terms 30 days',9)
    put(16,870,'SYNTHETIC ONLY - NOT A REAL INVOICE',8)
    if path.suffix=='.png':
        pix=page.get_pixmap(matrix=fitz.Matrix(2,2)); img=Image.open(io.BytesIO(pix.tobytes('png')))
        if layout==15: img=img.rotate(1.1,expand=True,fillcolor='white').filter(ImageFilter.GaussianBlur(.25))
        img.save(path);img.close()
    else: doc.save(path)
    doc.close()


class Diverse(base.Scenario):
    def __init__(self, idx, profile, root):
        name,sector,variant,layout=profile
        no=str(101+idx);lang='zh' if any(ord(c)>127 for c in name) else 'en'
        ccy={'fx':'EUR','btc':'BTC','eth':'ETH','usdt':'USDT'}.get(variant,'USD')
        fmt='xlsx' if layout in (8,9,18) else 'docx' if layout in (6,10,19) else 'png' if layout in (4,15) else 'pdf'
        super().__init__((no,name,lang,sector,ccy,fmt,variant),root)
        # Prevent dependence on a user's chart/classification configuration.
        config.CHART_PATH=self.root/'runtime/chart.json'; A.reload_chart()
        self.layout=layout; self.variant=variant;self.profit=D(0);self.cash=D(10000)
        self.result.update(layout_family=LAYOUTS[layout], checkpoints=[], expected_events=[])

    def correct(self,h,truth):
        for key,expected in truth.items():
            actual=db.get_invoice(h).f(key).value
            if key in ('subtotal','sales_tax','total_due'):
                try:
                    if D(str(actual))==D(expected):continue
                except Exception:pass
            if actual==expected:continue
            super().correct(h,{key:expected})

    def checkpoint(self,label):
        _,r=self.call('GET','/api/reports')
        self.result['checkpoints'].append({'stage':label,'report':r,'expected_profit':str(self.profit),'expected_cash':str(self.cash)})
        for k,v in r['checks'].items():
            if isinstance(v,dict): self.check(label+' '+k,v['ok'],v,category='ties')
        self.check(label+' 独立净利润',D(r['income_statement']['net_income'])==self.profit,r['income_statement']['net_income'],str(self.profit),'oracle')
        self.check(label+' 独立现金余额',D(r['cash_flow_statement']['ending'])==self.cash,r['cash_flow_statement']['ending'],str(self.cash),'oracle')
        _,v=self.call('GET','/api/ledger/open')
        self.check(label+' 往来明细与总账',all(x['ok'] for x in v['control'].values()),v['control'],category='control')
        return r

    def event(self,label,rows,activity=None,profit=0,cash=0):
        no=self.post('/api/ledger/manual',{'date':'2026-08-15','memo':label,'activity':activity,
            'allow_control':True,'counterparty':'Synthetic reviewed counterparty',
            'lines':[{'account':self.account(a),'debit':str(d),'credit':str(c)} for a,d,c in rows]},label)['entry_no']
        self.profit+=D(str(profit));self.cash+=D(str(cash))
        self.result['expected_events'].append({'label':label,'profit_delta':str(profit),'cash_delta':str(cash),'entry_no':no})
        self.checkpoint(label);return no

    def reverse(self,no,label,profit=0,cash=0):
        self.post('/api/ledger/reverse',{'entry_no':no},label)
        self.profit+=D(str(profit));self.cash+=D(str(cash));self.checkpoint(label)

    def settle(self,h,amount,book=None,activity=None,profit=0,cash=None,reason=None,ccy=None):
        label='结算 '+str(amount)
        body={'file_hash':h,'cash_amount':str(amount),'date':'2026-08-20'}
        if book is not None:body['settle_amount']=str(book)
        if activity:body['activity']=activity
        if reason:body['diff_reason']=reason
        if ccy:body['cash_currency']=ccy
        self.post('/api/ledger/settle',body,label)
        self.profit+=D(str(profit)); self.cash+=D(str(cash));self.checkpoint(label)

    def run(self):
        no,name,lang,sector,ccy,fmt,v=self.b
        self.post('/api/counterparties/register',{'name':name,'kind':'self'},'建档本方主体')
        self.post('/api/ledger/opening',{'date':'2026-07-31','other_lines':[
            {'account':self.account('1002'),'amount':'10000','side':'debit'},
            {'account':self.account('3000'),'amount':'10000','side':'credit'}]},'期初银行及资本')
        self.checkpoint('期初')
        asset=v in ('asset_unpaid','asset_partial','asset_paid','intangible','asset_credit')
        ap=asset or v in ('inventory','credit_ap')
        direction='AP' if ap else 'AR'
        credit=v in ('credit_ar','credit_ap','asset_credit')
        code='1601' if v in ('intangible','asset_credit') else '1500' if asset else '1405' if v=='inventory' else '6900' if ap else '4000'
        gross=D('-1100' if credit else '1100');net=D('-1000' if credit else '1000');tax=gross-net
        if ccy in ('BTC','ETH','USDT'):
            gross=D({'BTC':'0.01234567','ETH':'0.34567891','USDT':'1100.123456'}[ccy]);net=gross;tax=D(0)
        rate=D(base.RATES[ccy]);usd=(gross*rate).quantize(D('.01'));netusd=(net*rate).quantize(D('.01'))
        if ccy!='USD':self.post('/api/fx/rate',{'currency':ccy,'date':'2026-08-01','rate':str(rate)},'合成估值汇率')
        issuer='模拟供应商有限公司' if ap and lang=='zh' else 'Synthetic Equipment Vendor Ltd' if ap else name
        customer=name if ap else '模拟企业客户' if lang=='zh' else 'Synthetic Customer Ltd'
        truth=dict(invoice_no='DIVERSE-'+no,invoice_date='2026-08-05',issuer_name=issuer,customer_name=customer,
                   currency_settlement=ccy,subtotal=str(net),sales_tax=str(tax),total_due=str(gross))
        p=self.inputs/('source.'+fmt);fixture(p,self.layout,name,issuer,customer,truth)
        hashes=self.upload(p)
        self.check('单张票据不误拆分',len(hashes)==1,len(hashes),1,'extraction')
        if len(hashes)!=1:raise RuntimeError('Unexpected invoice splitting; evidence retained, no forced merge')
        h=hashes[0];inv=db.get_invoice(h)
        extracted={k:inv.f(k).value for k in truth}
        self.result['documents'].append({'file':str(p.relative_to(self.root)),'truth':truth,'extracted':extracted,
            'parse_method':inv.parse_method,'issues':[vars(i) for i in inv.issues],
            'line_items':[vars(i) for i in inv.line_items],
            'field_provenance':{k:{'raw':inv.f(k).raw,'source':inv.f(k).source} for k in truth}})
        for k,t in truth.items():
            a=extracted[k]
            try:ok=D(str(a))==D(t) if k in ('subtotal','sales_tax','total_due') else a==t
            except Exception:ok=False
            self.check('首次提取 '+k,ok,a,t,'extraction')
        self.correct(h,truth);self.approve(h,code)
        self.post('/api/ledger/post',{'file_hash':h,'direction':direction,'as_of':'2026-08-05'},'发票应计入账')
        if not ap:self.profit+=netusd
        elif not asset and v!='inventory':self.profit-=netusd
        self.checkpoint('发票入账未结算')
        if v in ('asset_partial','asset_paid','intangible'):
            self.settle(h,'440','440',cash=-440)
            if v!='asset_partial':self.settle(h,'660',cash=-660)
        elif v in ('credit_ar','credit_ap','asset_credit'):
            n=440 if v=='credit_ar' else 1100
            self.settle(h,n,n,cash=-n if v=='credit_ar' else n)
        elif v=='fx':
            self.post('/api/fx/rate',{'currency':'EUR','date':'2026-08-20','rate':'1.20'},'结算日汇率变动')
            self.settle(h,'1100',str(usd),profit=110,cash=1320,ccy='EUR',reason='fx_gain_loss')
        elif v=='fee':self.settle(h,1080,profit=-20,cash=1080,reason='fee')
        elif v=='withholding':self.settle(h,1000,cash=1000,reason='withholding_ar')
        elif v=='btc':self.settle(h,str(usd),cash=usd)
        elif v not in ('asset_unpaid','usdt','eth','inventory'):
            self.settle(h,550,550,cash=550)
        if v=='depreciation':
            self.event('付现购置投影设备',[('1500',2000,0),('1002',0,2000)],'investing',cash=-2000)
            e=self.event('计提折旧',[('6602',100,0),('1509',0,100)],profit=-100)
            self.reverse(e,'撤销误提折旧',profit=100)
            self.event('重提正确折旧',[('6602',80,0),('1509',0,80)],profit=-80)
        elif v=='disposal':
            self.event('购置设备',[('1500',1000,0),('1002',0,1000)],'investing',cash=-1000)
            self.event('计提折旧',[('6602',200,0),('1509',0,200)],profit=-200)
            e=self.event('设备处置',[('1002',900,0),('1509',200,0),('1500',0,1000),('4200',0,100)],'investing',profit=100,cash=900)
            self.reverse(e,'撤销处置',profit=-100,cash=-900)
        elif v=='intangible':
            e=self.event('无形资产摊销',[('6602',40,0),('1602',0,40)],profit=-40)
            self.reverse(e,'冲回摊销',profit=40)
        elif v=='inventory':
            self.event('已售部分结转成本',[('6401',400,0),('1405',0,400)],profit=-400)
            self.event('库存转自用设备',[('1500',300,0),('1405',0,300)])
        elif v=='prepaid':
            self.event('预付租金',[('1123',1200,0),('1002',0,1200)],'operating',cash=-1200)
            self.event('摊销首月租金',[('6610',100,0),('1123',0,100)],profit=-100)
        elif v=='advance':
            self.event('预收房款',[('1002',600,0),('2202',0,600)],'operating',cash=600)
            self.event('履约确认部分收入',[('2202',200,0),('4000',0,200)],profit=200)
        elif v=='payroll':
            self.event('计提工资',[('6601',300,0),('2241',0,300)],profit=-300)
            self.event('支付部分工资',[('2241',180,0),('1002',0,180)],'operating',cash=-180)
        elif v=='vat':
            self.event('销项转应交增值税',[('2210',100,0),('2211',0,100)])
            self.event('缴纳增值税',[('2211',100,0),('1002',0,100)],'operating',cash=-100)
        elif v=='loan':
            self.event('筹资借款暂挂其它应付',[('1002',3000,0),('2280',0,3000)],'financing',cash=3000)
            self.event('还本金',[('2280',800,0),('1002',0,800)],'financing',cash=-800)
        elif v=='novation':
            self.event('材料赊购',[('1405',700,0),('2280',0,700)])
            self.event('协议转长期借款',[('2280',700,0),('2501',0,700)])
            self.event('归还长期借款',[('2501',200,0),('1002',0,200)],'financing',cash=-200)
        elif v=='dividend':
            self.event('宣告股利',[('3200',200,0),('2280',0,200)])
            self.event('支付股利',[('2280',200,0),('1002',0,200)],'financing',cash=-200)
        elif v=='equity_pay':self.event('权益结算薪酬（测试凭证）',[('6601',150,0),('3000',0,150)],profit=-150)
        elif v=='credit_disposal':
            self.event('购置设备',[('1500',600,0),('1002',0,600)],'investing',cash=-600)
            self.event('赊销处置设备',[('1122',650,0),('1500',0,600),('4200',0,50)],profit=50)
            self.event('收回设备处置款',[('1002',650,0),('1122',0,650)],'investing',cash=650)
        elif v=='nonprofit':
            self.event('收到捐赠（模拟政策）',[('1002',500,0),('4200',0,500)],'operating',profit=500,cash=500)
            self.event('公益项目支出',[('6900',200,0),('1002',0,200)],'operating',profit=-200,cash=-200)
        elif v=='transfer':self.event('银行取现内部划转',[('1001',300,0),('1002',0,300)])
        elif v=='custody':
            self.event('代收客户本金',[('1002',5000,0),('2280',0,5000)],'operating',cash=5000)
            self.event('转出部分本金',[('2280',4500,0),('1002',0,4500)],'operating',cash=-4500)
        if v in ('usdt','eth'):
            status,data=self.call('POST','/api/ledger/settle',{'file_hash':h,'cash_amount':str(gross),'cash_currency':ccy,'settle_amount':str(usd)})
            self.check('代币原币结算明确拒绝',status==400 and '数字资产' in str(data),data,category='guard')
        self.post('/api/ledger/close',{'period':'2026-08'},'八月关账');self.checkpoint('八月关账后')
        if v=='periods':
            self.post('/api/ledger/manual',{'date':'2026-09-02','memo':'September service','activity':'operating','lines':[
                {'account':self.account('1002'),'debit':'300'},{'account':self.account('4000'),'credit':'300'}]},'九月服务收款')
            self.profit+=300;self.cash+=300;self.checkpoint('跨期新增业务')
            self.post('/api/ledger/close',{'period':'2026-09'},'九月关账');self.checkpoint('两期关账后')
            self.post('/api/ledger/reopen',{'period':'2026-09'},'重开九月');self.checkpoint('重开后')
            self.post('/api/ledger/close',{'period':'2026-09'},'九月再次关账');self.checkpoint('再次关账')
        r=self.checkpoint('最终');base.dump(self.root/'reports-after-close.json',r)
        status,result=self.call('POST','/api/reports/export',{})
        self.result['report_exported']=status==200 and 'download' in result
        self.check('最终报表导出',self.result['report_exported'],result,category='reports')
        if self.result['report_exported']:
            data=self.client.get(result['download']).content;(self.root/'financial-statements.xlsx').write_bytes(data)
            w=load_workbook(io.BytesIO(data),read_only=True);self.check('导出七张表',len(w.sheetnames)==7,category='reports');w.close()
        self.result['complete']=True


def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--ids',default='');args=p.parse_args()
    args.output.mkdir(parents=True,exist_ok=True);logging.basicConfig(level=logging.ERROR)
    base.dump(args.output/'manifest.json',{'profiles':PROFILES,'layout_families':LAYOUTS,'synthetic':True})
    out=[]
    for idx,profile in enumerate(PROFILES):
        no=str(101+idx)
        if args.ids and no not in args.ids.split(','):continue
        s=None;print('START',no,profile[0],flush=True)
        try:s=Diverse(idx,profile,args.output);s.run()
        except Exception:
            error=traceback.format_exc();print(error,flush=True)
            if s:s.result['fatal']=error
        finally:
            if s:
                out.append(s.finish());print('END',no,s.result['passed'],s.result['failed'],s.result['complete'],flush=True)
        base.dump(args.output/'summary.json',out)
    raise SystemExit(1 if any(s['failed'] or not s['complete'] for s in out) else 0)


if __name__=='__main__':main()
