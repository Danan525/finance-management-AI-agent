"""Structured headers, cover sheets and unit prices must not corrupt invoice ties."""
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal as D
from types import SimpleNamespace
from unittest.mock import patch
import fitz
from openpyxl import Workbook
from core import config, db
from extraction import pipeline
from extraction.parse import generic
from extraction.extract import excel
from extraction.extract import ocr


class DiverseLayouts(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        up=self.root/'uploads';up.mkdir()
        self.cfg=patch.multiple(config,DB_PATH=self.root/'test.db',UPLOAD_DIR=up,FX_AUTO_FETCH=False)
        self.cfg.start();db._initialized=False;db.init_db()

    def tearDown(self):
        self.cfg.stop();db._initialized=False;self.tmp.cleanup()

    def test_delimited_quantity_and_unit_price_not_duplicate_items(self):
        lines=['Description / Qty / Unit price / Amount',
               'Design work / 2 / 200.00 / 400.00',
               'Support work / 3 / 200.00 / 600.00', 'List price: 1100.00',
               'Commercial discount: 100.00', 'Subtotal: 1000.00']
        items=generic.extract_line_items([SimpleNamespace(words=[(0,500,s)]) for s in lines])
        self.assertEqual(len(items),2)
        self.assertEqual([i['quantity'] for i in items],['2','3'])
        self.assertEqual(sum(D(i['amount']) for i in items),D(1000))

    def test_chinese_invoice_title_and_net_label(self):
        doc=fitz.open();p=doc.new_page(width=720,height=900)
        lines=['服务发票','销售方名称: 模拟供应商有限公司','购买方名称: 模拟客户有限公司',
               '发票号码: CN-30','开票日期: 2026-08-05','Currency: USD',
               '未税金额: 1000.00','税额: 100.00','价税合计: 1100.00']
        for i,s in enumerate(lines):p.insert_text((30,40+i*35),s,fontname='china-s',fontsize=12)
        inv=pipeline.process_upload(doc.tobytes(),'chinese.pdf')[0];doc.close()
        self.assertEqual(inv.f('issuer_name').value,'模拟供应商有限公司')
        self.assertEqual(inv.f('subtotal').value,D(1000))

    def test_english_seller_label_not_saved_as_name(self):
        out=generic.extract_generic([SimpleNamespace(words=[(0,500,s)]) for s in
            ['Seller: New Coastal Services Ltd','Bill To: Client Ltd','Invoice No: EN-30']])
        self.assertEqual(out['issuer_name'],'New Coastal Services Ltd')
        self.assertFalse(generic.looks_like_company('Seller:'))
        self.assertFalse(generic.looks_like_company('Credit Note'))

    def test_excel_cover_sheet_does_not_hide_single_invoice(self):
        w=Workbook();w.active.title='Cover';w.active.append(['Instructions'])
        s=w.create_sheet('Invoice');s.append(['Invoice No','XLS-30']);s.append(['Total Due','1100.00'])
        path=self.root/'cover.xlsx';w.save(path);w.close()
        d=excel.excel_to_pdfdoc(path)
        self.assertIn('XLS-30',d.full_text);self.assertNotIn('Instructions',d.full_text)

    def test_excel_ambiguous_sheets_not_arbitrarily_selected(self):
        w=Workbook();w.active.append(['Cover remains ambiguous'])
        for n in ('A','B'):
            s=w.create_sheet(n);s.append(['Invoice No',n]);s.append(['Total Due','100.00'])
        path=self.root/'ambiguous.xlsx';w.save(path);w.close()
        self.assertIn('Cover remains ambiguous',excel.excel_to_pdfdoc(path).full_text)

    def test_horizontal_net_tax_total_not_mistaken_for_vat_items(self):
        lines=[SimpleNamespace(words=[(0,80,'未税金额'),(250,300,'税额'),(500,570,'价税合计')]),
               SimpleNamespace(words=[(0,80,'1000.00'),(250,300,'100.00'),(500,570,'1100.00')])]
        self.assertIsNone(generic.extract_cn_vat(lines))
        out=generic.extract_generic(lines)
        self.assertEqual([out[k] for k in ('subtotal','sales_tax','total_due')],['1000.00','100.00','1100.00'])

    def test_skew_reading_order_preserves_original_highlight_geometry(self):
        tokens=[]
        for text,x,base_y in [('Invoice No:',0,20),('INV-7',800,20),('Subtotal:',0,100),('100.00',800,100)]:
            y=base_y+.02*x
            box=[[x,y],[x+120,y+2.4],[x+120,y+22.4],[x,y+20]]
            tokens.append((text,.99,box))
        r=ocr.OcrResult(ocr._rows_to_text(tokens),tokens,.99,dpi=72,img_size=(1000,200))
        doc=r.to_pdfdoc()
        self.assertEqual(len(doc.lines),2)
        self.assertIn('Invoice No: INV-7',r.text)
        self.assertEqual(doc.words_geom[1][2],36)  # Original y, not straightened y.


if __name__=='__main__':unittest.main()
